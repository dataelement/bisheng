"""把纠错表中的管理员上传文件改挂到表里姓名对应的用户。

在 src/backend 执行。默认只生成预览,不写库:

  .venv/bin/python scripts/reassign_admin_upload_owners.py \\
    --tenant-id 1 --output /tmp/reassign-plan.jsonl

审核计划后写入。通过的行一次提交,跳过的行不写:

  .venv/bin/python scripts/reassign_admin_upload_owners.py \\
    --plan /tmp/reassign-plan.jsonl --apply --output /tmp/reassign-write.jsonl

回退先预览,确认后再加 --apply:

  .venv/bin/python scripts/reassign_admin_upload_owners.py \\
    --rollback /tmp/reassign-write.jsonl --output /tmp/reassign-rollback.jsonl

只改 knowledgefile 的 user_id、user_name、belonging_organization_id。
组织取目标用户当前唯一有效主部门,规则与 backfill_knowledge_file_organization.py 相同。
本脚本会覆盖已有所属组织。只改表格里的文件 ID,同一 document 的其它入口只记录。
不改 original_uploader_id、更新人、积分、权限和 Elasticsearch。
"""

from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import sys
import traceback
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

_BACKEND_ROOT = str(Path(__file__).resolve().parents[1])
if _BACKEND_ROOT not in sys.path:
    sys.path.insert(0, _BACKEND_ROOT)

from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import Column, Integer, MetaData, String, Table, update  # noqa: E402
from sqlmodel import col, select  # noqa: E402
from sqlmodel.ext.asyncio.session import AsyncSession  # noqa: E402

from bisheng.core.context.tenant import bypass_tenant_filter  # noqa: E402
from bisheng.database.models.department import Department  # noqa: E402
from bisheng.knowledge.domain.models.knowledge_file import KnowledgeFile  # noqa: E402
from bisheng.knowledge.domain.repositories.implementations.knowledge_file_organization_repository_impl import (  # noqa: E402
    KnowledgeFileOrganizationRepositoryImpl,
)

ADMIN_UPLOADER_NAME = "管理员"
DEFAULT_XLSX = Path(__file__).resolve().parent / "纠错2.xlsx"
# The user domain model pulls the auth import chain. This script only reads three columns.
_USER_LOOKUP = Table(
    "user",
    MetaData(),
    Column("user_id", Integer),
    Column("user_name", String),
    Column("delete", Integer, key="deleted_flag"),
)
SHEET_COLUMNS = ("关联文件入口ID", "文档唯一标识", "文件名", "上传人ID")
QUERY_CHUNK = 500


@dataclass(frozen=True)
class SheetRow:
    """校验通过、且文件 ID 不冲突的一行纠错数据。"""

    excel_row: int
    file_id: int
    document_id: int
    file_name: str
    target_name: str


class JsonlJournal:
    """写前持久化日志,不覆盖已有报告。"""

    def __init__(self, path: str | Path):
        fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
        self.file = os.fdopen(fd, "w", encoding="utf-8")

    def write(self, value: dict) -> None:
        self.file.write(json.dumps(value, ensure_ascii=False, sort_keys=True, default=str) + "\n")
        self.file.flush()
        os.fsync(self.file.fileno())

    def __enter__(self):
        return self

    def __exit__(self, *_args):
        self.file.close()


def database_key(session: AsyncSession) -> str:
    url = session.get_bind().url
    identity = (url.drivername, url.host, url.port, url.database)
    return hashlib.sha256(json.dumps(identity).encode()).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_positive_int(value: Any) -> int | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float) and value.is_integer():
        number = int(value)
        return number if number > 0 else None
    text = str(value).strip()
    if not text.isdigit():
        return None
    number = int(text)
    return number if number > 0 else None


def parse_document_id(value: Any) -> int | None:
    text = "" if value is None else str(value).strip()
    prefix = "document:"
    if not text.startswith(prefix):
        return None
    return parse_positive_int(text[len(prefix) :])


def as_datetime(value: Any) -> datetime | None:
    if value is None or isinstance(value, datetime):
        return value
    return datetime.fromisoformat(str(value).replace(" ", "T"))


def dump_time(value: Any) -> str | None:
    parsed = as_datetime(value)
    if parsed is None:
        return None
    return parsed.isoformat()


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def classify_named_users(rows: list[tuple[int, str, int]]) -> tuple[dict[str, Any] | None, str | None]:
    """同一姓名的用户行。恰好一个 delete=0 才可用。"""

    active = [row for row in rows if int(row[2] or 0) == 0]
    if len(active) == 1:
        user_id, user_name, _deleted = active[0]
        return {"user_id": int(user_id), "user_name": user_name}, None
    if len(active) > 1:
        return None, "user_name_ambiguous"
    if rows:
        return None, "user_resigned"
    return None, "user_not_found"


def sheet_skip(excel_row: int, file_id: int | None, target_name: str, reason: str, **extra: Any) -> dict:
    record = {
        "type": "skipped",
        "excel_row": excel_row,
        "file_id": file_id,
        "target_name": target_name,
        "reason": reason,
    }
    record.update(extra)
    return record


def load_sheet(path: Path) -> tuple[list[SheetRow], list[dict], list[dict], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("表格没有表头")
        indexes: dict[str, int] = {}
        for position, name in enumerate(header):
            label = cell_text(name)
            if label in SHEET_COLUMNS and label in indexes:
                raise ValueError(f"重复列: {label}")
            if label:
                indexes[label] = position
        missing = [name for name in SHEET_COLUMNS if name not in indexes]
        if missing:
            raise ValueError("缺少列: " + ", ".join(missing))

        valid: list[SheetRow] = []
        skipped: list[dict] = []
        data_rows = 0
        for excel_row, values in enumerate(rows, start=2):
            if values is None or all(cell is None or cell_text(cell) == "" for cell in values):
                continue
            data_rows += 1
            raw = [values[indexes[name]] if indexes[name] < len(values) else None for name in SHEET_COLUMNS]
            file_id = parse_positive_int(raw[0])
            document_id = parse_document_id(raw[1])
            file_name = cell_text(raw[2])
            target_name = cell_text(raw[3])
            if file_id is None:
                skipped.append(sheet_skip(excel_row, None, target_name, "sheet_file_id_invalid"))
            elif document_id is None:
                skipped.append(sheet_skip(excel_row, file_id, target_name, "sheet_document_invalid"))
            elif not file_name:
                skipped.append(sheet_skip(excel_row, file_id, target_name, "sheet_file_name_missing"))
            elif not target_name:
                skipped.append(sheet_skip(excel_row, file_id, target_name, "sheet_target_name_missing"))
            else:
                valid.append(
                    SheetRow(
                        excel_row=excel_row,
                        file_id=file_id,
                        document_id=document_id,
                        file_name=file_name,
                        target_name=target_name,
                    )
                )
    finally:
        workbook.close()

    grouped: dict[int, list[SheetRow]] = defaultdict(list)
    for row in valid:
        grouped[row.file_id].append(row)
    unique: list[SheetRow] = []
    notes: list[dict] = []
    for file_id, items in grouped.items():
        identities = {(item.document_id, item.file_name, item.target_name) for item in items}
        if len(identities) > 1:
            for item in items:
                skipped.append(sheet_skip(item.excel_row, file_id, item.target_name, "sheet_conflict"))
            continue
        items.sort(key=lambda item: item.excel_row)
        unique.append(items[0])
        if len(items) > 1:
            notes.append(
                {
                    "type": "duplicate_collapsed",
                    "file_id": file_id,
                    "excel_rows": [item.excel_row for item in items],
                }
            )
    unique.sort(key=lambda item: item.excel_row)
    return unique, skipped, notes, data_rows


def file_problem(row: SheetRow, record: dict | None, tenant_id: int) -> dict | None:
    if record is None:
        return sheet_skip(row.excel_row, row.file_id, row.target_name, "file_not_found")
    if int(record["tenant_id"]) != tenant_id:
        return sheet_skip(row.excel_row, row.file_id, row.target_name, "cross_tenant")
    if cell_text(record["file_name"]) != row.file_name:
        return sheet_skip(
            row.excel_row,
            row.file_id,
            row.target_name,
            "file_name_mismatch",
            db_file_name=record["file_name"],
        )
    if record["reference_document_id"] != row.document_id:
        return sheet_skip(
            row.excel_row,
            row.file_id,
            row.target_name,
            "document_identity_mismatch",
            db_document_id=record["reference_document_id"],
        )
    if cell_text(record["user_name"]) != ADMIN_UPLOADER_NAME:
        return sheet_skip(
            row.excel_row,
            row.file_id,
            row.target_name,
            "uploader_not_admin",
            db_user_name=record["user_name"],
        )
    return None


def owner_snapshot(record: dict) -> dict:
    return {
        "user_id": record["user_id"],
        "user_name": record["user_name"],
        "belonging_organization_id": record["belonging_organization_id"],
        "update_time": dump_time(record["update_time"]),
    }


def same_owner(record: dict, expected: dict) -> bool:
    return (
        record["user_id"] == expected["user_id"]
        and record["user_name"] == expected["user_name"]
        and record["belonging_organization_id"] == expected["belonging_organization_id"]
    )


def file_columns():
    return (
        KnowledgeFile.id,
        KnowledgeFile.tenant_id,
        KnowledgeFile.user_id,
        KnowledgeFile.user_name,
        KnowledgeFile.file_name,
        KnowledgeFile.reference_document_id,
        KnowledgeFile.belonging_organization_id,
        KnowledgeFile.update_time,
    )


def file_record(row: Any) -> dict:
    return {
        "id": int(row.id),
        "tenant_id": int(row.tenant_id),
        "user_id": None if row.user_id is None else int(row.user_id),
        "user_name": row.user_name,
        "file_name": row.file_name,
        "reference_document_id": None if row.reference_document_id is None else int(row.reference_document_id),
        "belonging_organization_id": None
        if row.belonging_organization_id is None
        else int(row.belonging_organization_id),
        "update_time": as_datetime(row.update_time),
    }


async def files_by_ids(session: AsyncSession, file_ids: list[int]) -> dict[int, dict]:
    found: dict[int, dict] = {}
    for start in range(0, len(file_ids), QUERY_CHUNK):
        chunk = file_ids[start : start + QUERY_CHUNK]
        rows = (await session.exec(select(*file_columns()).where(col(KnowledgeFile.id).in_(chunk)))).all()
        for row in rows:
            found[int(row.id)] = file_record(row)
    return found


async def users_by_names(session: AsyncSession, names: list[str]) -> dict[str, list[tuple[int, str, int]]]:
    found: dict[str, list[tuple[int, str, int]]] = {name: [] for name in names}
    for start in range(0, len(names), QUERY_CHUNK):
        chunk = names[start : start + QUERY_CHUNK]
        rows = (
            await session.exec(
                select(
                    _USER_LOOKUP.c.user_id,
                    _USER_LOOKUP.c.user_name,
                    _USER_LOOKUP.c.deleted_flag,
                ).where(_USER_LOOKUP.c.user_name.in_(chunk))
            )
        ).all()
        for row in rows:
            found.setdefault(row.user_name, []).append((int(row.user_id), row.user_name, int(row.deleted_flag or 0)))
    return found


async def organizations_for(session: AsyncSession, user_ids: list[int], tenant_id: int) -> dict[int, int | None]:
    repository = KnowledgeFileOrganizationRepositoryImpl(session)
    return {user_id: await repository.primary_organization(user_id, tenant_id) for user_id in user_ids}


async def department_names(session: AsyncSession, department_ids: list[int]) -> dict[int, str]:
    if not department_ids:
        return {}
    rows = (
        await session.exec(select(Department.id, Department.name).where(col(Department.id).in_(department_ids)))
    ).all()
    return {int(row.id): row.name for row in rows}


async def admin_siblings(session: AsyncSession, document_ids: list[int], sheet_file_ids: set[int]) -> list[dict]:
    siblings: list[dict] = []
    columns = (
        KnowledgeFile.id,
        KnowledgeFile.reference_document_id,
        KnowledgeFile.file_name,
        KnowledgeFile.user_id,
        KnowledgeFile.user_name,
        KnowledgeFile.entry_type,
        KnowledgeFile.deleted_at,
    )
    for start in range(0, len(document_ids), QUERY_CHUNK):
        chunk = document_ids[start : start + QUERY_CHUNK]
        rows = (await session.exec(select(*columns).where(col(KnowledgeFile.reference_document_id).in_(chunk)))).all()
        for row in rows:
            if int(row.id) in sheet_file_ids or cell_text(row.user_name) != ADMIN_UPLOADER_NAME:
                continue
            siblings.append(
                {
                    "type": "sibling",
                    "file_id": int(row.id),
                    "document_id": int(row.reference_document_id),
                    "file_name": row.file_name,
                    "user_id": None if row.user_id is None else int(row.user_id),
                    "user_name": row.user_name,
                    "entry_type": row.entry_type,
                    "deleted": row.deleted_at is not None,
                }
            )
    siblings.sort(key=lambda item: int(item["file_id"]))
    return siblings


def build_plan(
    rows: list[SheetRow],
    sheet_skips: list[dict],
    files: dict[int, dict],
    users: dict[str, list[tuple[int, str, int]]],
    organizations: dict[int, int | None],
    names: dict[int, str],
    tenant_id: int,
) -> tuple[list[dict], list[dict]]:
    candidates: list[dict] = []
    skipped = list(sheet_skips)
    resolved_users: dict[str, tuple[dict[str, Any] | None, str | None]] = {}
    for row in rows:
        problem = file_problem(row, files.get(row.file_id), tenant_id)
        if problem is not None:
            skipped.append(problem)
            continue
        if row.target_name not in resolved_users:
            resolved_users[row.target_name] = classify_named_users(users.get(row.target_name, []))
        user, reason = resolved_users[row.target_name]
        if user is None:
            skipped.append(sheet_skip(row.excel_row, row.file_id, row.target_name, reason or "user_not_found"))
            continue
        organization_id = organizations.get(int(user["user_id"]))
        if organization_id is None:
            skipped.append(
                sheet_skip(
                    row.excel_row,
                    row.file_id,
                    row.target_name,
                    "uploader_no_unique_valid_primary_organization",
                    user_id=user["user_id"],
                )
            )
            continue
        current = files[row.file_id]
        candidates.append(
            {
                "type": "candidate",
                "excel_row": row.excel_row,
                "file_id": row.file_id,
                "document_id": row.document_id,
                "file_name": current["file_name"],
                "tenant_id": tenant_id,
                "target_name": row.target_name,
                "user_id": user["user_id"],
                "user_name": user["user_name"],
                "organization_id": organization_id,
                "organization_name": names.get(organization_id),
                "before": owner_snapshot(current),
            }
        )
    skipped.sort(key=lambda item: (item.get("excel_row") or 0, item.get("file_id") or 0))
    return candidates, skipped


async def preview(
    session: AsyncSession,
    journal: JsonlJournal,
    *,
    tenant_id: int,
    xlsx_path: Path,
) -> dict:
    if tenant_id <= 0:
        raise ValueError("tenant_id必须为正整数")
    rows, sheet_skips, notes, data_rows = load_sheet(xlsx_path)
    with bypass_tenant_filter():
        files = await files_by_ids(session, [row.file_id for row in rows])
        names = sorted({row.target_name for row in rows})
        user_rows = await users_by_names(session, names)
        resolved_ids = []
        for name in names:
            user, _reason = classify_named_users(user_rows.get(name, []))
            if user is not None:
                resolved_ids.append(int(user["user_id"]))
        organizations = await organizations_for(session, resolved_ids, tenant_id)
        org_names = await department_names(session, [org_id for org_id in organizations.values() if org_id is not None])
        document_ids = sorted({row.document_id for row in rows})
        siblings = await admin_siblings(session, document_ids, {row.file_id for row in rows})
        await session.rollback()
    candidates, skipped = build_plan(rows, sheet_skips, files, user_rows, organizations, org_names, tenant_id)
    summary = {
        "type": "summary",
        "sheet_rows": data_rows,
        "would_update": len(candidates),
        "skipped": len(skipped),
        "siblings": len(siblings),
        "skipped_by_reason": dict(Counter(item["reason"] for item in skipped)),
    }
    journal.write(
        {
            "type": "header",
            "version": 1,
            "mode": "preview",
            "tenant_id": tenant_id,
            "database_key": database_key(session),
            "xlsx_sha256": file_sha256(xlsx_path),
            "rule": "sheet_uploader_name_current_primary_organization",
        }
    )
    for record in (*candidates, *skipped, *notes, *siblings, summary):
        journal.write(record)
    return {
        "would_update": summary["would_update"],
        "skipped": summary["skipped"],
        "siblings": summary["siblings"],
    }


def read_records(path: str | Path):
    with open(path, encoding="utf-8") as stream:
        for line in stream:
            if not line.endswith("\n"):
                yield {"type": "incomplete_tail"}
                return
            record = json.loads(line)
            if not isinstance(record, dict):
                raise ValueError("报告记录必须为JSON对象")
            yield record


def load_complete_report(path: str | Path) -> list[dict]:
    records = list(read_records(path))
    if not records or records[-1].get("type") == "incomplete_tail":
        raise ValueError("报告不完整,请重新生成")
    if records[0].get("type") != "header" or records[-1].get("type") != "summary":
        raise ValueError("报告缺少表头或汇总")
    return records


def require_header(header: dict, session: AsyncSession, mode: str) -> int:
    tenant_id = header.get("tenant_id", 0)
    if (
        header.get("version") != 1
        or header.get("mode") != mode
        or header.get("database_key") != database_key(session)
        or not isinstance(tenant_id, int)
        or tenant_id <= 0
    ):
        raise ValueError("报告版本、模式、租户或数据库身份不匹配")
    return tenant_id


async def lock_file(session: AsyncSession, file_id: int) -> dict | None:
    row = (await session.exec(select(*file_columns()).where(KnowledgeFile.id == file_id).with_for_update())).first()
    if row is None:
        return None
    return file_record(row)


def values_match_plan(record: dict | None, candidate: dict) -> bool:
    before = candidate["before"]
    return (
        record is not None
        and int(record["tenant_id"]) == int(candidate["tenant_id"])
        and record["file_name"] == candidate["file_name"]
        and record["reference_document_id"] == candidate["document_id"]
        and same_owner(record, before)
    )


async def planned_owner_still_valid(session: AsyncSession, candidate: dict) -> bool:
    rows = (
        await session.exec(
            select(
                _USER_LOOKUP.c.user_id,
                _USER_LOOKUP.c.user_name,
                _USER_LOOKUP.c.deleted_flag,
            ).where(_USER_LOOKUP.c.user_id == candidate["user_id"])
        )
    ).all()
    if len(rows) != 1 or int(rows[0].deleted_flag or 0) != 0 or rows[0].user_name != candidate["user_name"]:
        return False
    organization_id = await KnowledgeFileOrganizationRepositoryImpl(session).primary_organization(
        int(candidate["user_id"]),
        int(candidate["tenant_id"]),
    )
    return organization_id == candidate["organization_id"]


async def write_owner(session: AsyncSession, candidate: dict, owner: dict, expected: dict) -> bool:
    predicates = [
        KnowledgeFile.id == candidate["file_id"],
        KnowledgeFile.tenant_id == candidate["tenant_id"],
        KnowledgeFile.file_name == candidate["file_name"],
        KnowledgeFile.reference_document_id == candidate["document_id"],
        KnowledgeFile.user_name == expected["user_name"],
    ]
    if expected["user_id"] is None:
        predicates.append(KnowledgeFile.user_id.is_(None))
    else:
        predicates.append(KnowledgeFile.user_id == expected["user_id"])
    if expected["belonging_organization_id"] is None:
        predicates.append(KnowledgeFile.belonging_organization_id.is_(None))
    else:
        predicates.append(KnowledgeFile.belonging_organization_id == expected["belonging_organization_id"])
    _result = await session.exec(
        update(KnowledgeFile)
        .where(*predicates)
        .values(
            user_id=owner["user_id"],
            user_name=owner["user_name"],
            belonging_organization_id=owner["belonging_organization_id"],
            update_time=as_datetime(owner.get("update_time", expected.get("update_time"))),
        )
    )
    current = await files_by_ids(session, [int(candidate["file_id"])])
    written = current.get(int(candidate["file_id"]))
    return written is not None and same_owner(written, owner)


def after_owner(candidate: dict) -> dict:
    return {
        "user_id": candidate["user_id"],
        "user_name": candidate["user_name"],
        "belonging_organization_id": candidate["organization_id"],
    }


async def apply_plan(session: AsyncSession, path: Path, journal: JsonlJournal) -> dict:
    records = load_complete_report(path)
    tenant_id = require_header(records[0], session, "preview")
    candidates = [record for record in records if record.get("type") == "candidate"]
    committed: list[dict] = []
    skipped = 0
    with bypass_tenant_filter():
        journal.write(
            {
                "type": "header",
                "version": 1,
                "mode": "apply",
                "tenant_id": tenant_id,
                "database_key": database_key(session),
                "source_plan": str(path),
                "rule": "sheet_uploader_name_current_primary_organization",
            }
        )
        document_ids = sorted({int(item["document_id"]) for item in candidates})
        sheet_ids = {
            int(item["file_id"])
            for item in records
            if item.get("type") in {"candidate", "skipped"} and item.get("file_id")
        }
        for sibling in await admin_siblings(session, document_ids, sheet_ids):
            journal.write(sibling)
        try:
            for candidate in candidates:
                if int(candidate.get("tenant_id", 0)) != tenant_id:
                    raise ValueError("计划中的租户与报告表头不一致")
                current = await lock_file(session, int(candidate["file_id"]))
                if not values_match_plan(current, candidate) or not await planned_owner_still_valid(session, candidate):
                    skipped += 1
                    journal.write(
                        {
                            "type": "skipped",
                            "file_id": candidate["file_id"],
                            "excel_row": candidate.get("excel_row"),
                            "reason": "plan_state_changed",
                        }
                    )
                    continue
                owner = after_owner(candidate)
                changed = await write_owner(session, candidate, owner, candidate["before"])
                if not changed:
                    skipped += 1
                    journal.write(
                        {
                            "type": "skipped",
                            "file_id": candidate["file_id"],
                            "excel_row": candidate.get("excel_row"),
                            "reason": "plan_state_changed",
                        }
                    )
                    continue
                committed.append({"file_id": candidate["file_id"], "before": candidate["before"], "after": owner})
            if not committed:
                await session.rollback()
            else:
                journal.write({"type": "prepared", "files": committed})
                await session.commit()
                journal.write({"type": "committed", "files": committed})
        except Exception:
            await session.rollback()
            raise
    summary = {"type": "summary", "updated": len(committed), "skipped": skipped}
    journal.write(summary)
    return {"updated": len(committed), "skipped": skipped}


async def rollback_plan(session: AsyncSession, path: Path, journal: JsonlJournal, *, apply: bool) -> dict:
    records = load_complete_report(path)
    tenant_id = require_header(records[0], session, "apply")
    prepared = next((record for record in records if record.get("type") == "prepared"), None)
    committed = next((record for record in records if record.get("type") == "committed"), None)
    mode = "rollback" if apply else "rollback-preview"
    journal.write(
        {
            "type": "header",
            "version": 1,
            "mode": mode,
            "tenant_id": tenant_id,
            "database_key": database_key(session),
            "rule": "sheet_uploader_name_current_primary_organization",
        }
    )
    if prepared is not None and committed is None:
        journal.write({"type": "skipped", "reason": "commit_outcome_unknown"})
        summary = {"type": "summary", "updated": 0, "skipped": 1, "would_update": 0}
        journal.write(summary)
        return {"updated": 0, "skipped": 1, "would_update": 0}
    files = list(committed.get("files", [])) if committed else []
    restored: list[dict] = []
    skipped = 0
    would_update = 0
    with bypass_tenant_filter():
        try:
            for item in files:
                current = await lock_file(session, int(item["file_id"]))
                if current is None or int(current["tenant_id"]) != tenant_id or not same_owner(current, item["after"]):
                    skipped += 1
                    journal.write({"type": "skipped", "file_id": item["file_id"], "reason": "rollback_state_changed"})
                    continue
                would_update += 1
                if not apply:
                    journal.write({"type": "would_rollback", "file_id": item["file_id"]})
                    continue
                changed = await write_owner(
                    session,
                    {
                        "file_id": item["file_id"],
                        "tenant_id": tenant_id,
                        "file_name": current["file_name"],
                        "document_id": current["reference_document_id"],
                    },
                    item["before"],
                    item["after"],
                )
                if not changed:
                    skipped += 1
                    journal.write({"type": "skipped", "file_id": item["file_id"], "reason": "rollback_state_changed"})
                    continue
                restored.append(item)
            if not apply or not restored:
                await session.rollback()
            else:
                journal.write({"type": "prepared", "files": restored})
                await session.commit()
                journal.write({"type": "committed", "files": restored})
        except Exception:
            await session.rollback()
            raise
    summary = {"type": "summary", "updated": len(restored), "skipped": skipped, "would_update": would_update}
    journal.write(summary)
    return {"updated": len(restored), "skipped": skipped, "would_update": would_update}


def build_parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    result.add_argument("--tenant-id", type=int, help="预览时必须指定租户")
    result.add_argument("--xlsx", type=Path, help=f"纠错表路径, 默认 {DEFAULT_XLSX.name}")
    mode = result.add_mutually_exclusive_group()
    mode.add_argument("--plan", type=Path, help="执行已审核的预览计划,需要 --apply")
    mode.add_argument("--rollback", type=Path, help="按写入日志回退,默认只预览")
    result.add_argument("--apply", action="store_true", help="实际写入")
    result.add_argument("--output", type=Path, required=True, help="新报告路径,不允许覆盖")
    return result


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = build_parser()
    args = parser.parse_args(argv)
    if args.plan or args.rollback:
        if args.tenant_id is not None or args.xlsx is not None:
            parser.error("执行计划或回退时不能再指定租户或表格")
        if args.plan and not args.apply:
            parser.error("执行计划必须显式传入 --apply")
    elif args.tenant_id is None or args.apply:
        parser.error("预览必须指定 --tenant-id;写入需先生成计划,再使用 --plan 和 --apply")
    elif args.xlsx is None:
        args.xlsx = DEFAULT_XLSX
    return args


async def run(args: argparse.Namespace) -> int:
    from bisheng.core.context.manager import close_app_context
    from bisheng.core.database import get_async_db_session

    try:
        with JsonlJournal(args.output) as journal:
            async with get_async_db_session() as session:
                if args.plan:
                    counts = await apply_plan(session, args.plan, journal)
                elif args.rollback:
                    counts = await rollback_plan(session, args.rollback, journal, apply=args.apply)
                else:
                    counts = await preview(session, journal, tenant_id=args.tenant_id, xlsx_path=args.xlsx)
        print(json.dumps(counts, ensure_ascii=False, sort_keys=True))
        wrote = bool(args.apply)
        return 3 if wrote and counts.get("skipped") else 0
    finally:
        await close_app_context()


def main(argv: list[str] | None = None) -> int:
    try:
        args = parse_args(argv)
    except SystemExit as exc:
        code = exc.code
        return code if isinstance(code, int) else 2
    try:
        return asyncio.run(run(args))
    except (ValueError, FileExistsError, FileNotFoundError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    except KeyboardInterrupt:
        return 130
    except Exception as exc:
        print(f"执行失败({type(exc).__name__}),请保留报告并核对提交结果", file=sys.stderr)
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())
