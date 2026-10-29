"""Office samples named by the release plan. Opening them proves structure, not desktop Office."""

from __future__ import annotations

import io

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation


def xlsx_amounts() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "金额"
    sheet["A1"] = "金额"
    for index, amount in enumerate((1, 2, 3), start=2):
        sheet.cell(row=index, column=1, value=amount)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def docx_body(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def pptx_body(text: str) -> bytes:
    presentation = Presentation()
    layout = presentation.slide_layouts[5]
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = text
    buf = io.BytesIO()
    presentation.save(buf)
    return buf.getvalue()


def xlsx_contains_amounts(payload: bytes, amounts: tuple[int, ...]) -> bool:
    book = load_workbook(io.BytesIO(payload), data_only=True)
    seen: list[int] = []
    for sheet in book.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if isinstance(cell, (int, float)) and int(cell) == cell:
                    seen.append(int(cell))
    return all(amount in seen for amount in amounts)


def docx_contains(payload: bytes, text: str) -> bool:
    document = Document(io.BytesIO(payload))
    return any(text in (paragraph.text or "") for paragraph in document.paragraphs)


def pptx_contains(payload: bytes, text: str) -> bool:
    presentation = Presentation(io.BytesIO(payload))
    for slide in presentation.slides:
        for shape in slide.shapes:
            if text in getattr(shape, "text", ""):
                return True
    return False
