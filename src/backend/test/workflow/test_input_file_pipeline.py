"""Workflow extraction keeps images without applying retrieval chunk constraints."""

import shutil
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock

import openpyxl
import pytest
from langchain_core.documents import Document
from PIL import Image

from bisheng.common.errcode.knowledge import KnowledgeFileChunkMaxError
from bisheng.knowledge.rag.pipeline.loader.base import BaseBishengLoader
from bisheng.knowledge.rag.pipeline.loader.x_create import XinChuangFormatterLoader
from bisheng.knowledge.rag.pipeline.types import ImageStorageContext
from bisheng.knowledge.rag.temp_file_pipeline import TempFilePipeline
from bisheng.workflow.nodes.input.input import InputNode


class _ImageLoader(BaseBishengLoader):
    def load(self):
        image = Path(self.ensure_local_image_dir()) / "figure.png"
        image.write_bytes(b"image-bytes")
        text = f"Before\n![figure]({self.build_image_url(image.name)})\nAfter"
        return [
            Document(
                page_content=text,
                metadata={
                    "indexes": [[0, len(text)]],
                    "pages": [1],
                    "types": ["text"],
                    "bboxes": [[0, 0, 10, 10]],
                },
            )
        ]


def _pipeline(tmp_path, filename="file.pdf", **kwargs):
    path = tmp_path / filename
    path.touch()
    pipeline = TempFilePipeline(1, str(path), filename, **kwargs)
    pipeline.__dict__["file_metadata"] = {}
    return pipeline


@pytest.mark.parametrize("split", [False, True])
def test_image_upload_precedes_staging_cleanup_and_preserves_offsets(tmp_path, monkeypatch, split):
    """AC-01, AC-04: bytes survive staging cleanup and offsets use the final image URL."""
    objects = {}
    events = []
    storage = MagicMock()
    storage.put_object_sync.side_effect = lambda **kw: objects.update(
        {kw["object_name"]: Path(kw["file"]).read_bytes()}
    )
    monkeypatch.setattr(
        "bisheng.knowledge.rag.pipeline.transformer.image_upload.get_minio_storage_sync", lambda: storage
    )
    image_context = ImageStorageContext(
        "chat/1/images",
        "main",
        lambda name: f"/api/v1/chat/c/files/{name}/content",
        before_upload=lambda names: events.append(("pending", names.copy())),
        after_upload=lambda: events.append(("ready", bool(objects))),
    )
    pipeline = _pipeline(tmp_path, split_for_retrieval=split, image_storage=image_context)
    monkeypatch.setattr(pipeline, "_init_pdf_loader", lambda: _ImageLoader(**pipeline._get_loader_common_params()))
    result = pipeline.run()
    assert objects == {"chat/1/images/figure.png": b"image-bytes"}
    assert events == [("pending", ["figure.png"]), ("ready", True)]
    assert not Path(pipeline.loader.local_image_dir).exists()
    text = "\n".join(doc.page_content for doc in result.documents)
    assert "/api/v1/chat/c/files/figure.png/content" in text
    assert pipeline.tmp_dir not in text
    if split:
        assert result.documents[0].metadata["chunk_bboxes"][0]["page"] == 1
    else:
        assert result.documents[0].metadata["indexes"] == [[0, len(text)]]
    assert storage.put_object_sync.call_args.kwargs["bucket_name"] == "main"


def test_failed_manifest_finalize_removes_uploaded_bytes(tmp_path, monkeypatch):
    """AC-08: a deleted conversation cannot leave an upload completed after deletion."""
    storage = MagicMock()
    monkeypatch.setattr(
        "bisheng.knowledge.rag.pipeline.transformer.image_upload.get_minio_storage_sync", lambda: storage
    )

    def fail():
        raise ValueError("conversation deleted")

    context = ImageStorageContext(
        "chat/images", "main", lambda name: f"/image/{name}", after_upload=fail, rollback_on_failure=True
    )
    pipeline = _pipeline(tmp_path, image_storage=context, split_for_retrieval=False)
    monkeypatch.setattr(pipeline, "_init_pdf_loader", lambda: _ImageLoader(**pipeline._get_loader_common_params()))
    with pytest.raises(ValueError, match="conversation deleted"):
        pipeline.run()
    storage.remove_object_sync.assert_called_once_with(bucket_name="main", object_name="chat/images/figure.png")


def test_partial_upload_never_marks_manifest_ready_and_cleans_completed_objects(tmp_path, monkeypatch):
    """AC-08: failure on the second image cannot publish a partially uploaded result."""

    class TwoImages(_ImageLoader):
        def load(self):
            documents = super().load()
            (Path(self.local_image_dir) / "second.png").write_bytes(b"second-image")
            return documents

    storage = MagicMock()
    storage.put_object_sync.side_effect = [None, OSError("upload interrupted")]
    monkeypatch.setattr(
        "bisheng.knowledge.rag.pipeline.transformer.image_upload.get_minio_storage_sync", lambda: storage
    )
    ready = MagicMock()
    context = ImageStorageContext(
        "chat/images", "main", lambda name: f"/image/{name}", after_upload=ready, rollback_on_failure=True
    )
    pipeline = _pipeline(tmp_path, image_storage=context, split_for_retrieval=False)
    monkeypatch.setattr(pipeline, "_init_pdf_loader", lambda: TwoImages(**pipeline._get_loader_common_params()))
    with pytest.raises(OSError, match="upload interrupted"):
        pipeline.run()
    ready.assert_not_called()
    uploaded = storage.put_object_sync.call_args_list[0].kwargs["object_name"]
    storage.remove_object_sync.assert_called_once_with(bucket_name="main", object_name=uploaded)


def test_unrecorded_manifest_does_not_upload_bytes(tmp_path, monkeypatch):
    """AC-08: a DB registration failure cannot create untracked permanent objects."""
    storage = MagicMock()
    monkeypatch.setattr(
        "bisheng.knowledge.rag.pipeline.transformer.image_upload.get_minio_storage_sync", lambda: storage
    )

    def fail(_names):
        raise OSError("manifest unavailable")

    context = ImageStorageContext("chat/images", "main", lambda name: f"/image/{name}", before_upload=fail)
    pipeline = _pipeline(tmp_path, image_storage=context, split_for_retrieval=False)
    monkeypatch.setattr(pipeline, "_init_pdf_loader", lambda: _ImageLoader(**pipeline._get_loader_common_params()))
    with pytest.raises(OSError, match="manifest unavailable"):
        pipeline.run()
    storage.put_object_sync.assert_not_called()


@pytest.mark.parametrize("split", [False, True])
def test_ppt_long_page_is_checked_only_for_retrieval(tmp_path, monkeypatch, split):
    """AC-02, AC-09: extracting a long slide succeeds; the existing ingest limit stays."""
    pipeline = _pipeline(tmp_path, "slides.pptx", split_for_retrieval=split)
    monkeypatch.setattr(
        pipeline,
        "_init_ppt_loader",
        lambda: SimpleNamespace(
            load=lambda: [Document(page_content="x" * 10001, metadata={"page": 1})],
        ),
    )
    if split:
        with pytest.raises(KnowledgeFileChunkMaxError):
            pipeline.run()
    else:
        assert pipeline.run().documents[0].page_content == "x" * 10001


def test_extract_excel_keeps_one_table_per_sheet_and_long_cell(tmp_path):
    """AC-02, AC-03: no row chunks, duplicate headers, or retrieval size budget."""
    pipeline = _pipeline(tmp_path, "wide.xlsx", split_for_retrieval=False)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["unique-header", "value"])
    sheet.append(["x" * 12000, "long"])
    for i in range(25):
        sheet.append([str(i), "row"])
    workbook.save(pipeline.local_file_path)
    documents = pipeline.run().documents
    assert len(documents) == 1
    assert "x" * 12000 in documents[0].page_content
    assert documents[0].page_content.count("unique-header") == 1
    assert "| 24 | row |" in documents[0].page_content


def test_et_delegate_receives_extraction_and_image_context(tmp_path):
    """AC-03, AC-09: conversion cannot accidentally restore Excel retrieval splitting."""

    def url_builder(name):
        return f"/file/{name}"

    loader = XinChuangFormatterLoader(
        file_path="input.et",
        file_metadata={},
        file_extension="et",
        tmp_dir=str(tmp_path),
        split_for_retrieval=False,
        image_object_dir="chat/images",
        image_url_builder=url_builder,
    )
    delegate = loader._build_delegate_loader("converted.xlsx", "xlsx", "excel")
    assert delegate.split_for_retrieval is False
    assert delegate.image_url_builder is url_builder
    assert delegate.image_object_dir == "chat/images"


@pytest.mark.parametrize("mode,split", [("extract_text", False), (["extract_text", "ingest_to_temp_kb"], True)])
def test_input_node_selects_purpose_before_parsing_and_keeps_output_limit(monkeypatch, mode, split):
    """AC-03, AC-09: parser purpose is per file, and extraction still caps the body."""
    node = InputNode.__new__(InputNode)
    node.id, node.workflow_id = "input", "workflow"
    node._image_ext = ["png", "jpg"]
    node._vector_client, node._es_client = MagicMock(), MagicMock()
    node.init_vector_clients = MagicMock()
    node.get_upload_file_path_content = MagicMock(return_value=(["x" * 20000], [{}]))
    monkeypatch.setattr(
        "bisheng.workflow.nodes.input.input.KnowledgeService.get_upload_file_original_name", lambda name: name
    )
    result = node.parse_upload_file(
        "k",
        {
            "key": "k",
            "file_content": "c",
            "file_path": "p",
            "image_file": "img",
            "file_parse_mode": mode,
            "file_content_size": 15000,
        },
        ["http://storage/temp/file.pdf"],
    )
    node.get_upload_file_path_content.assert_called_once_with("http://storage/temp/file.pdf", split_for_retrieval=split)
    assert "x" * 15000 in result["c"] and "x" * 15001 not in result["c"]
    assert node.init_vector_clients.call_count == int(split)
    assert node._vector_client.add_texts.call_count == int(split)
    assert node._es_client.add_texts.call_count == int(split)


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx", "pdf"])
def test_real_document_images_reach_object_storage_before_cleanup(tmp_path, monkeypatch, extension):
    """AC-01: real local loaders emit stable references backed by uploaded image bytes."""
    if extension == "docx" and not shutil.which("pandoc"):
        pytest.skip("local DOCX image extraction requires pandoc")
    image_path = tmp_path / "source.png"
    Image.new("RGB", (16, 16), (240, 10, 30)).save(image_path)
    pipeline = _pipeline(tmp_path, f"source.{extension}", split_for_retrieval=False)
    if extension == "docx":
        from docx import Document as WordDocument

        document = WordDocument()
        document.add_paragraph("Document with an image")
        document.add_picture(str(image_path))
        document.save(pipeline.local_file_path)
    elif extension == "pptx":
        from pptx import Presentation
        from pptx.util import Inches

        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
        presentation.save(pipeline.local_file_path)
        monkeypatch.setattr("bisheng.knowledge.rag.pipeline.loader.ppt.convert_ppt_to_pdf", lambda *args: None)
    elif extension == "xlsx":
        from openpyxl.drawing.image import Image as ExcelImage

        workbook = openpyxl.Workbook()
        workbook.active.append(["header", "value"])
        workbook.active.add_image(ExcelImage(str(image_path)), "A3")
        workbook.save(pipeline.local_file_path)
    else:
        import fitz

        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((30, 30), "Document with an image")
            page.insert_image(fitz.Rect(50, 50, 200, 200), filename=str(image_path))
            document.save(pipeline.local_file_path)
        # Pin the local PDF loader; provider HTTP clients are outside this test.
        from bisheng.knowledge.rag.pipeline.loader.pdf import LocalPdfLoader

        monkeypatch.setattr(
            pipeline,
            "_init_pdf_loader",
            lambda: LocalPdfLoader(
                **pipeline._get_loader_common_params(),
                retain_images=True,
            ),
        )

    objects = {}
    storage = MagicMock()
    storage.put_object_sync.side_effect = lambda **kw: objects.update(
        {kw["object_name"]: Path(kw["file"]).read_bytes()}
    )
    monkeypatch.setattr(
        "bisheng.knowledge.rag.pipeline.transformer.image_upload.get_minio_storage_sync", lambda: storage
    )
    pipeline.image_storage = ImageStorageContext("chat/images", "main", lambda name: f"/stable-images/{name}")
    result = pipeline.run()
    assert objects, f"{extension} did not stage and upload its embedded image"
    text = "\n".join(doc.page_content for doc in result.documents)
    for name, content in objects.items():
        assert content
        assert f"/stable-images/{name.rsplit('/', 1)[1]}" in text
    assert not Path(pipeline.tmp_dir).exists()
