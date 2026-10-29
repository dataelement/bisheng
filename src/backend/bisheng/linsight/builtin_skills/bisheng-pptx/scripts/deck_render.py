#!/usr/bin/env python
"""Native chart / table rendering shared by the template generator and ``Deck``.

Styles come from the template's ``layouts.json`` (``chart_style`` / ``style`` on
the slot), so a chart built into a deck looks like the template sample.
Library only — import it, don't run it.
"""

from __future__ import annotations

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Pt

__all__ = ["CHART_KINDS", "add_chart", "build_table", "fill_table", "set_run_font"]

CHART_KINDS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, latin: str, ea: str) -> None:
    """Write <a:latin> and <a:ea>; <a:ea> must immediately follow <a:latin>."""
    rPr = run._r.get_or_add_rPr()
    latin_el = rPr.get_or_add_latin()
    latin_el.set("typeface", latin)
    ea_el = rPr.find(qn("a:ea"))
    if ea_el is None:
        ea_el = latin_el.makeelement(qn("a:ea"), {})
        latin_el.addnext(ea_el)
    ea_el.set("typeface", ea)


# ---------------------------------------------------------------------------
# Charts
# ---------------------------------------------------------------------------


def _waterfall_data(categories, values):
    """Start, deltas…, end → invisible base series + visible bar series.

    ``values[0]`` is the starting total, ``values[-1]`` the ending total, the
    rest are signed contributions. Negative contributions hang down from the
    running total. Stacked columns cannot float a bar below zero, so a running
    total that drops under zero is rejected.
    """
    base, bars = [], []
    running = 0.0
    last = len(values) - 1
    for i, v in enumerate(values):
        if i == 0 or i == last:
            base.append(0.0)
            bars.append(abs(v))
            running = v
        else:
            low = running + min(v, 0.0)
            if low < 0:
                raise ValueError("waterfall 的累计值不能跌破 0；这组数据请改用 column 或 bar")
            base.append(low)
            bars.append(abs(v))
            running += v
    return base, bars


def add_chart(shapes, x, y, w, h, kind, categories, series, style, number_format="0"):
    """Add a styled native chart at (x, y, w, h) points and return its graphic frame.

    ``series`` is a list of ``(label, values)``. For ``waterfall`` pass a single
    series whose first and last values are totals and the middle values are deltas.
    """
    if kind not in CHART_KINDS:
        raise ValueError(f"unknown chart kind {kind!r}; use one of {sorted(CHART_KINDS)}")
    data = CategoryChartData()
    data.categories = list(categories)
    if kind == "waterfall":
        label, values = series[0]
        base, bars = _waterfall_data(categories, list(values))
        data.add_series("_base", base)
        data.add_series(label, bars)
    else:
        for label, values in series:
            data.add_series(label, list(values))
    frame = shapes.add_chart(CHART_KINDS[kind], Pt(x), Pt(y), Pt(w), Pt(h), data)
    _style_chart(frame.chart, kind, style, number_format, series)
    return frame


def _style_chart(chart, kind, style, number_format, series):
    palette = style["palette"]
    chart.font.size = Pt(11)
    chart.font.color.rgb = _rgb(style["muted"])
    chart.font.name = style["font"]
    plot = chart.plots[0]
    visible_series = len(series)
    if kind in ("pie", "doughnut"):
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.RIGHT
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        points = plot.series[0].points
        for i in range(len(plot.categories)):
            fill = points[i].format.fill
            fill.solid()
            fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        plot.has_data_labels = True
        plot.data_labels.number_format = number_format
        plot.data_labels.number_format_is_linked = False
        plot.data_labels.font.size = Pt(11)
        plot.data_labels.font.color.rgb = _rgb(style["text"])
        return

    chart.has_legend = visible_series > 1 and kind != "waterfall"
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
    if hasattr(plot, "gap_width"):
        plot.gap_width = 60 if kind == "waterfall" else 80
    if kind in ("stacked", "waterfall"):
        plot.overlap = 100
    for i, s in enumerate(plot.series):
        fmt = s.format
        if kind == "waterfall" and i == 0:
            fmt.fill.background()
            fmt.line.fill.background()
            continue
        color = palette[(i - 1 if kind == "waterfall" else i) % len(palette)]
        if kind == "line":
            fmt.line.color.rgb = _rgb(color)
            fmt.line.width = Pt(2.25)
        else:
            fmt.fill.solid()
            fmt.fill.fore_color.rgb = _rgb(color)
        s_labels = s.data_labels
        s_labels.show_value = True
        s_labels.number_format = number_format
        s_labels.number_format_is_linked = False
        s_labels.font.size = Pt(10)
        s_labels.font.color.rgb = _rgb(style["text"])
        if kind in ("column", "bar", "line"):
            s_labels.position = XL_LABEL_POSITION.OUTSIDE_END if kind != "line" else XL_LABEL_POSITION.ABOVE
    if kind == "waterfall":  # colour totals vs. deltas
        bars = plot.series[1]
        last = len(plot.categories) - 1
        values = list(series[0][1])
        for i in range(len(plot.categories)):
            fill = bars.points[i].format.fill
            fill.solid()
            if i in (0, last):
                fill.fore_color.rgb = _rgb(palette[0])
            else:
                fill.fore_color.rgb = _rgb(palette[1] if values[i] >= 0 else style.get("negative", palette[2]))
                if values[i] < 0:  # the bar stores abs(v); put the sign back on its label
                    _point_label(bars.points[i], "-" + _format_number(abs(values[i]), number_format), style)
    for axis in (chart.category_axis, chart.value_axis):
        axis.format.line.color.rgb = _rgb(style["line"])
        axis.tick_labels.font.size = Pt(10)
        axis.tick_labels.font.color.rgb = _rgb(style["muted"])
    chart.value_axis.has_major_gridlines = True
    chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(style["line"])
    chart.value_axis.format.line.fill.background()


def _format_number(value: float, number_format: str) -> str:
    """Python rendering of the simple Excel formats used here (``0``, ``0.0``, ``#,##0.00``)."""
    decimals = len(number_format.split(".", 1)[1].rstrip("%")) if "." in number_format else 0
    text = f"{value:,.{decimals}f}" if "," in number_format else f"{value:.{decimals}f}"
    return text + "%" if number_format.endswith("%") else text


def _point_label(point, text: str, style) -> None:
    label = point.data_label
    label.has_text_frame = True
    label.text_frame.text = text
    run = label.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(10)
    run.font.color.rgb = _rgb(style["text"])
    label.position = XL_LABEL_POSITION.OUTSIDE_END


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


def _cell_text(cell, text, size, color, bold, style):
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(8)
    tf.margin_top = tf.margin_bottom = Pt(4)
    tf.text = ""
    run = tf.paragraphs[0].add_run()
    run.text = "" if text is None else str(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    set_run_font(run, style["latin"], style["font"])
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _cell_bottom_border(cell, color: str, width_pt: float = 0.75) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    fills = [f for tag in ("a:solidFill", "a:noFill") for f in tcPr.findall(qn(tag))]
    for f in fills:
        tcPr.remove(f)
    for tag, visible in (("a:lnL", False), ("a:lnR", False), ("a:lnT", False), ("a:lnB", True)):
        ln = etree.SubElement(tcPr, qn(tag), {"w": str(int(width_pt * 12700)) if visible else "0"})
        if visible:
            fill = etree.SubElement(ln, qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), {"val": color})
        else:
            etree.SubElement(ln, qn("a:noFill"))
    for f in fills:  # cell fill comes after the borders in CT_TableCellProperties
        tcPr.append(f)


def fill_table(table, header, rows, style, col_widths=None, total_w=None, row_h=34.0, highlight_rows=()):
    """Style an existing python-pptx table with ``header`` + ``rows``."""
    tblPr = table._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    style_id = tblPr.find(qn("a:tableStyleId"))
    if style_id is None:
        style_id = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style_id.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
    if col_widths and total_w:
        scale = total_w / sum(col_widths)
        for col, cw in zip(table.columns, col_widths, strict=False):
            col.width = Pt(cw * scale)
    for row in table.rows:
        row.height = Pt(row_h)
    for c, text in enumerate(header):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = _rgb(style["header_fill"])
        _cell_text(cell, text, style["header_size"], style["header_color"], True, style)
    for r, values in enumerate(rows, start=1):
        hot = (r - 1) in highlight_rows
        for c in range(len(header)):
            text = values[c] if c < len(values) else ""
            cell = table.cell(r, c)
            cell.fill.solid()
            if hot:
                fill = style.get("highlight_fill", style["stripe_fill"])
            else:
                fill = style["stripe_fill"] if r % 2 == 0 else style["body_fill"]
            cell.fill.fore_color.rgb = _rgb(fill)
            color = style.get("highlight_color", style["body_color"]) if hot else style["body_color"]
            _cell_text(cell, text, style["body_size"], color, hot, style)
            _cell_bottom_border(cell, style["line_color"])


def build_table(shapes, x, y, w, header, rows, style, col_widths=None, row_h=34.0, highlight_rows=()):
    """Add a styled native table at (x, y) points, ``w`` wide; returns the graphic frame."""
    frame = shapes.add_table(len(rows) + 1, len(header), Pt(x), Pt(y), Pt(w), Pt(row_h * (len(rows) + 1)))
    fill_table(
        frame.table, header, rows, style, col_widths=col_widths, total_w=w, row_h=row_h, highlight_rows=highlight_rows
    )
    return frame
