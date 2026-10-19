"""在线课堂导入流转: 真实接口写入 171 MySQL, 再读表核对."""

from __future__ import annotations

import uuid
from contextlib import asynccontextmanager
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI
from fastapi.responses import JSONResponse
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import inspect as sa_inspect
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine
from sqlmodel.ext.asyncio.session import AsyncSession

from bisheng.common.dependencies.user_deps import UserPayload
from bisheng.common.errcode.base import BaseErrorCode
from bisheng.core.context.tenant import set_current_tenant_id
from bisheng.core.database.connection import _patch_aiomysql_pre_ping
from bisheng.shougang_portal_course.api.endpoints.admin import router as course_admin_router
from bisheng.shougang_portal_course.api.endpoints.catalog_admin import router as catalog_admin_router
from bisheng.shougang_portal_course.domain.services.yunxuetang_course_sheet import (
    SHEET_NAME,
    TEMPLATE_HEADERS,
)
from test.qa_expert.conftest import _mysql_async_url


def _workbook(rows: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        sheet.cell(row=1, column=col, value=header)
    for row_idx, row in enumerate(rows, start=2):
        for col, header in enumerate(TEMPLATE_HEADERS, start=1):
            if header in row:
                sheet.cell(row=row_idx, column=col, value=row[header])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _app(user) -> FastAPI:
    app = FastAPI()

    @app.exception_handler(BaseErrorCode)
    async def _biz(_request, exc: BaseErrorCode):
        return JSONResponse(
            {"status_code": exc.code, "status_message": exc.message, "data": None},
            status_code=200,
        )

    app.include_router(course_admin_router)
    app.include_router(catalog_admin_router)
    app.dependency_overrides[UserPayload.get_admin_user] = lambda: user
    return app


async def _cleanup(engine, prefix: str) -> None:
    async with AsyncSession(engine, expire_on_commit=False) as session:
        await session.execute(
            text("DELETE FROM portal_course WHERE external_id LIKE :p OR name LIKE :p"),
            {"p": f"{prefix}%"},
        )
        while True:
            rows = (
                await session.execute(
                    text("SELECT id, parent_id FROM portal_course_catalog WHERE name LIKE :p"),
                    {"p": f"{prefix}%"},
                )
            ).fetchall()
            if not rows:
                break
            parent_ids = {row[1] for row in rows if row[1]}
            leaves = [row[0] for row in rows if row[0] not in parent_ids]
            if not leaves:
                break
            for leaf_id in leaves:
                await session.execute(
                    text("DELETE FROM portal_course_catalog WHERE id = :id"),
                    {"id": leaf_id},
                )
        await session.commit()


async def _course_row(engine, external_id: str):
    async with AsyncSession(engine, expire_on_commit=False) as session:
        result = await session.execute(
            text(
                "SELECT name, instructor, description, tags_json, duration_seconds, "
                "external_url, enabled, show_on_home, sort_order, organization, cover_url, "
                "catalog_id, source_updated_at "
                "FROM portal_course WHERE external_id = :external_id"
            ),
            {"external_id": external_id},
        )
        return result.fetchall()


async def _catalog_rows(engine, prefix: str):
    async with AsyncSession(engine, expire_on_commit=False) as session:
        result = await session.execute(
            text(
                "SELECT name, parent_id, opened, order_index, catalog_name_path "
                "FROM portal_course_catalog WHERE name LIKE :p ORDER BY name"
            ),
            {"p": f"{prefix}%"},
        )
        return result.fetchall()


@pytest.fixture
async def mysql_engine():
    _patch_aiomysql_pre_ping()
    engine = create_async_engine(_mysql_async_url(), pool_pre_ping=True)
    async with engine.connect() as connection:
        ready = await connection.run_sync(
            lambda sync: "duration_seconds"
            in {item["name"] for item in sa_inspect(sync).get_columns("portal_course")}
        )
    if not ready:
        await engine.dispose()
        raise RuntimeError("portal_course.duration_seconds 不存在, 先执行 f117 迁移")
    yield engine
    await engine.dispose()


async def test_course_and_catalog_import_write_the_same_rows(mysql_engine, monkeypatch):
    prefix = f"yxflow{uuid.uuid4().hex[:8]}"
    external_id = f"{prefix}1"
    root, middle, leaf = f"{prefix}根", f"{prefix}中", f"{prefix}叶"
    path = f"{root}->{middle}->{leaf}"
    set_current_tenant_id(1)
    await _cleanup(mysql_engine, prefix)

    @asynccontextmanager
    async def patched_session():
        session = AsyncSession(bind=mysql_engine, expire_on_commit=False)
        try:
            yield session
        finally:
            await session.close()

    monkeypatch.setattr(
        "bisheng.shougang_portal_course.api.endpoints.admin.get_async_db_session",
        patched_session,
    )
    monkeypatch.setattr(
        "bisheng.shougang_portal_course.api.endpoints.catalog_admin.get_async_db_session",
        patched_session,
    )
    user = SimpleNamespace(user_id=7, tenant_id=1)
    first = _workbook(
        [
            {
                "课件名称": f"{prefix}课",
                "全路径": path,
                "分类名称": leaf,
                "主讲老师": "--",
                "标准时长": "30分0秒",
                "平台课程编号": external_id,
                "分享链接": "https://short-phx-ali.yunxuetang.cn/flow1",
                "课程简介": "第一次简介",
                "更新日期": "2026-09-11 09:06:08",
                "标签": "--",
            },
            {
                "课件名称": f"{prefix}坏行",
                "全路径": path,
                "分类名称": "对不上",
                "主讲老师": "吴铮",
                "标准时长": "10分0秒",
                "平台课程编号": f"{prefix}bad",
                "分享链接": "https://short-phx-ali.yunxuetang.cn/bad",
                "课程简介": "",
                "更新日期": "2026-09-11 09:06:08",
                "标签": "",
            },
        ]
    )
    second = _workbook(
        [
            {
                "课件名称": f"{prefix}课更新",
                "全路径": path,
                "分类名称": leaf,
                "主讲老师": "吴铮",
                "标准时长": "45分0秒",
                "平台课程编号": external_id,
                "分享链接": "https://short-phx-ali.yunxuetang.cn/flow2",
                "课程简介": "第二次简介",
                "更新日期": "2026-09-12 15:45:04",
                "标签": "采购",
            }
        ]
    )
    try:
        transport = ASGITransport(app=_app(user))
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            preview = await client.post(
                "/shougang-portal/course-admin/courses/import/preview",
                files={"file": ("在线课堂.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert preview.status_code == 200
            preview_body = preview.json()
            assert preview_body["status_code"] == 200
            assert preview_body["data"]["total"] == 2
            assert preview_body["data"]["valid"] == 1
            assert preview_body["data"]["catalog_count"] == 3
            assert preview_body["data"]["issues"]

            imported = await client.post(
                "/shougang-portal/course-admin/courses/import",
                files={"file": ("在线课堂.xlsx", first, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert imported.status_code == 200
            imported_body = imported.json()
            assert imported_body["status_code"] == 200
            assert imported_body["data"]["success"] == 1
            assert imported_body["data"]["failed"] == 1

        rows = await _course_row(mysql_engine, external_id)
        assert len(rows) == 1
        course = rows[0]
        assert course.name == f"{prefix}课"
        assert course.instructor == ""
        assert course.description == "第一次简介"
        assert course.tags_json == "[]"
        assert int(course.duration_seconds) == 1800
        assert course.external_url == "https://short-phx-ali.yunxuetang.cn/flow1"
        assert int(course.enabled) == 1
        assert int(course.show_on_home) == 1
        assert course.cover_url == ""
        assert str(course.source_updated_at).startswith("2026-09-11 09:06:08")
        bad_rows = await _course_row(mysql_engine, f"{prefix}bad")
        assert bad_rows == []
        catalogs = await _catalog_rows(mysql_engine, prefix)
        assert {item.name for item in catalogs} == {root, middle, leaf}
        leaf_row = next(item for item in catalogs if item.name == leaf)
        assert leaf_row.catalog_name_path == path
        assert int(leaf_row.opened) == 1
        assert int(leaf_row.order_index) == 0

        async with AsyncSession(mysql_engine, expire_on_commit=False) as session:
            await session.execute(
                text(
                    "UPDATE portal_course SET cover_url = :cover, show_on_home = 0, "
                    "sort_order = 6, organization = :org WHERE external_id = :external_id"
                ),
                {
                    "cover": "https://example.com/cover.jpg",
                    "org": "首钢股份",
                    "external_id": external_id,
                },
            )
            await session.execute(
                text(
                    "UPDATE portal_course_catalog SET opened = 0, order_index = 9 "
                    "WHERE id = :catalog_id"
                ),
                {"catalog_id": course.catalog_id},
            )
            await session.commit()

        async with AsyncClient(transport=transport, base_url="http://test") as client:
            again = await client.post(
                "/shougang-portal/course-admin/catalogs/import",
                files={"file": ("在线课堂.xlsx", second, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            )
            assert again.status_code == 200
            again_body = again.json()
            assert again_body["status_code"] == 200
            assert again_body["data"]["success"] == 1
            assert again_body["data"]["failed"] == 0

        rows = await _course_row(mysql_engine, external_id)
        assert len(rows) == 1
        updated = rows[0]
        assert updated.name == f"{prefix}课更新"
        assert updated.instructor == "吴铮"
        assert updated.description == "第二次简介"
        assert "采购" in updated.tags_json
        assert int(updated.duration_seconds) == 2700
        assert updated.external_url == "https://short-phx-ali.yunxuetang.cn/flow2"
        assert updated.cover_url == "https://example.com/cover.jpg"
        assert int(updated.show_on_home) == 0
        assert int(updated.sort_order) == 6
        assert updated.organization == "首钢股份"
        assert updated.catalog_id == course.catalog_id
        catalogs = await _catalog_rows(mysql_engine, prefix)
        assert len(catalogs) == 3
        leaf_row = next(item for item in catalogs if item.name == leaf)
        assert int(leaf_row.opened) == 0
        assert int(leaf_row.order_index) == 9
    finally:
        await _cleanup(mysql_engine, prefix)
