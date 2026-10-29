#!/usr/bin/env python
"""Template-aware checks used by ``inspect_deck.py`` (library, not a CLI).

Four checks, all reported as WARN until calibrated on real task output:

* **capacity** — a slot holds more 字数 than its template capacity;
* **leftover sample** — template sample text ("示例…") is still on the page, or a
  slot was never filled and still shows the template's sample value;
* **repetition** — three or more consecutive pages use the same layout
  (``layout_id`` from the manifest, otherwise the page's shape silhouette);
* **cover focus** — the cover has no first-read focal point: the largest text is
  under 3× the median text size and no picture/chart/table/filled shape covers
  ≥ 25% of the canvas.

With a manifest (written by ``Deck.save``) the existing inspect checks are also
diffed against the prototype page each slide was cloned from, so the template's
own structure (rules, badges, small footer text) is not reported as a defect.
"""

from __future__ import annotations

import json
import os
import re
import statistics

from deck_template import TEMPLATES_DIR, text_units
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

__all__ = ["apply_baseline", "find_manifest", "load_manifest", "template_checks"]

REPEAT_RUN = 3
COVER_RATIO = 3.0
COVER_VISUAL_SHARE = 0.25
SAMPLE_MARK = "示例"
# Template defaults that may legitimately stay as they are.
_KEEPABLE = {"目录", "CONTENTS", "感谢聆听", "培训课程", "学习目标", "本节内容", "高", "中", "低"}
_ORDINAL = re.compile(r"^(0?\d{1,2}|Q\d)$")


# --------------------------------------------------------------------------- #
# manifest
# --------------------------------------------------------------------------- #
def find_manifest(pptx_path: str, explicit: str | None = None) -> str | None:
    """``--manifest`` wins; otherwise ``scratch/<stem>.manifest.json`` next to cwd or the deck's parent."""
    if explicit:
        return explicit if os.path.exists(explicit) else None
    stem = os.path.splitext(os.path.basename(pptx_path))[0]
    candidates = [
        os.path.join("scratch", f"{stem}.manifest.json"),
        os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(pptx_path))), "scratch", f"{stem}.manifest.json"),
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


MANIFEST_SHAPE_MATCH = 0.6


def load_manifest(path: str, prs) -> tuple[dict | None, str]:
    """Return (manifest, note). A manifest that no longer describes ``prs`` is ignored.

    The page count must match, and when the manifest records each page's shape
    names (``Deck.save`` does), most of them must still be on that page: a deck
    rebuilt from scratch under the same file name keeps the count but not the names.
    """
    try:
        with open(path, encoding="utf-8") as fh:
            manifest = json.load(fh)
        spec_path = os.path.join(TEMPLATES_DIR, manifest["template"], "layouts.json")
        with open(spec_path, encoding="utf-8") as fh:
            manifest["_layouts"] = {item["layout_id"]: item for item in json.load(fh)["layouts"]}
    except Exception as exc:  # a broken manifest must not kill the report
        return None, f"模板清单 {path} 读取失败（{type(exc).__name__}），按从零生成的文稿检查"
    slides = list(prs.slides)
    if len(manifest["pages"]) != len(slides):
        return None, (
            f"模板清单 {path} 记录 {len(manifest['pages'])} 页，文件实际 {len(slides)} 页"
            f"（保存后又改过页数？）→ 忽略清单，按从零生成的文稿检查"
        )
    recorded = total = 0
    for page, slide in zip(manifest["pages"], slides, strict=True):
        names = set(page.get("shapes") or [])
        recorded += len(names)
        total += len(names & {shape.name for shape in slide.shapes})
    if recorded and total / recorded < MANIFEST_SHAPE_MATCH:
        return None, f"模板清单 {path} 和文件的页面对不上（文稿可能已重新生成）→ 忽略清单，按从零生成的文稿检查"
    return manifest, f"模板清单：{path}（模板 {manifest['template']}）"


# --------------------------------------------------------------------------- #
# baseline diff for the existing inspect checks
# --------------------------------------------------------------------------- #
def _finding_key(finding) -> tuple:
    # The whole message, numbers included: a template's own structure (rules,
    # badges, footer text) yields the same message on every clone, while a filled
    # slot that overflows by more than the sample did yields a different one.
    return finding.level, finding.where, finding.message


def apply_baseline(findings_by_slide: dict, manifest: dict, check_slide, slide_w: float, slide_h: float) -> int:
    """Drop findings that the prototype page itself produces. Returns how many were suppressed."""
    prs = Presentation(os.path.join(TEMPLATES_DIR, manifest["template"], "template.pptx"))
    cache: dict[str, set] = {}
    suppressed = 0
    for page in manifest["pages"]:
        layout = manifest["_layouts"].get(page["layout_id"])
        if layout is None:
            continue
        if page["layout_id"] not in cache:
            proto = prs.slides[layout["slide_index"]]
            cache[page["layout_id"]] = {_finding_key(f) for f in check_slide(proto, page["index"], slide_w, slide_h)}
        baseline = cache[page["layout_id"]]
        kept = []
        for f in findings_by_slide.get(page["index"], []):
            if _finding_key(f) in baseline:
                suppressed += 1
            else:
                kept.append(f)
        findings_by_slide[page["index"]] = kept
    return suppressed


# --------------------------------------------------------------------------- #
# the four checks
# --------------------------------------------------------------------------- #
def _shapes_by_name(slide) -> dict:
    return {shape.name: shape for shape in slide.shapes}


def _texts(slide):
    for shape in slide.shapes:
        if shape.has_text_frame:
            yield shape.name, shape.text_frame.text
        elif getattr(shape, "has_table", False) and shape.has_table:
            for r, row in enumerate(shape.table.rows):
                for c, cell in enumerate(row.cells):
                    yield f"{shape.name}[{r}][{c}]", cell.text


def _check_capacity_and_samples(prs, manifest, Finding) -> list:
    findings = []
    for index, slide in enumerate(prs.slides, start=1):
        for where, text in _texts(slide):
            if SAMPLE_MARK in text:
                findings.append(
                    Finding(
                        "WARN",
                        index,
                        where,
                        f"残留模板示例文字：{text.strip()[:30]!r} → 换成真实内容，用不到就 deck.drop",
                    )
                )
        if manifest is None:
            continue
        page = manifest["pages"][index - 1]
        layout = manifest["_layouts"].get(page["layout_id"])
        if layout is None:
            continue
        shapes = _shapes_by_name(slide)
        for slot in layout["slots"]:
            shape = shapes.get(slot["name"])
            if "capacity" in slot and shape is not None and shape.has_text_frame:
                used = text_units(shape.text_frame.text)
                if used > slot["capacity"]:
                    findings.append(
                        Finding(
                            "WARN",
                            index,
                            slot["name"],
                            f"{used:g} 字超过该位置建议容量 {slot['capacity']} 字 → 精简文案或拆页，不要缩字号",
                        )
                    )
        for name in page.get("unfilled", []):
            slot = next((s for s in layout["slots"] if s["name"] == name), None)
            if slot is None:
                continue
            if slot["role"] == "chart_area":
                findings.append(
                    Finding("WARN", index, name, "图表区没有放图表 → deck.chart(...)，或换一个不带图表的版式")
                )
                continue
            sample = slot.get("sample", "")
            if not sample or SAMPLE_MARK in sample or sample in _KEEPABLE or _ORDINAL.match(sample):
                continue  # 示例 text is reported above; ordinals and headings may stay
            if name in shapes:
                findings.append(
                    Finding("WARN", index, name, f"插槽没填写，仍是模板示例值 {sample!r} → 填真实内容或 deck.drop")
                )
    return findings


def _silhouette(slide, slide_w: float, slide_h: float) -> tuple:
    boxes = []
    for shape in slide.shapes:
        w, h = (shape.width or 0) / 12700, (shape.height or 0) / 12700
        if w * h < 400 or (w >= slide_w - 2 and h >= slide_h - 2):
            continue
        boxes.append(
            (round((shape.left or 0) / 12700 / 48), round((shape.top or 0) / 12700 / 48), round(w / 48), round(h / 48))
        )
    return tuple(sorted(boxes))


def _check_repetition(prs, manifest, Finding, slide_w, slide_h) -> list:
    if manifest is not None:
        keys = [page["layout_id"] for page in manifest["pages"]]
    else:
        keys = [_silhouette(slide, slide_w, slide_h) for slide in prs.slides]
    findings = []
    start = 0
    for i in range(1, len(keys) + 1):
        if i < len(keys) and keys[i] == keys[start]:
            continue
        run = i - start
        if run >= REPEAT_RUN and keys[start]:
            label = f"版式 {keys[start]}" if manifest is not None else "相同构图"
            findings.append(
                Finding(
                    "WARN",
                    start + 1,
                    "",
                    f"第 {start + 1}–{i} 页连续 {run} 页使用{label} → 换一种版式或调整主视觉，避免通篇一个样",
                )
            )
        start = i
    return findings


def _check_cover(prs, Finding, size_of, slide_w, slide_h) -> list:
    if not len(prs.slides):
        return []
    cover = prs.slides[0]
    canvas = slide_w * slide_h
    for shape in cover.shapes:
        w, h = (shape.width or 0) / 12700, (shape.height or 0) / 12700
        share = w * h / canvas if canvas else 0
        if share < COVER_VISUAL_SHARE or share >= 0.9:
            continue
        is_media = (
            shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.TABLE)
            or getattr(shape, "has_chart", False)
            or getattr(shape, "has_table", False)
        )
        filled = shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and shape.fill.type == 1  # MSO_FILL.SOLID
        if is_media or filled:
            return []
    sizes = []
    for shape in cover.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            biggest = max(
                (
                    size_of(shape, para, run)
                    for para in shape.text_frame.paragraphs
                    for run in para.runs
                    if run.text.strip()
                ),
                default=0,
            )
            if biggest:
                sizes.append(biggest)
    if len(sizes) < 2:
        return []
    largest, median = max(sizes), statistics.median(sizes)
    if median and largest / median >= COVER_RATIO:
        return []
    return [
        Finding(
            "WARN",
            1,
            "",
            f"封面缺少首读焦点：最大字号 {largest:.0f}pt 只有文字字号中位数 {median:.0f}pt 的 "
            f"{largest / median:.1f} 倍 → 加大主标题（≥3 倍），或放一张占画布四分之一以上的图",
        )
    ]


def template_checks(prs, manifest, Finding, size_of, slide_w: float, slide_h: float) -> list:
    """All four new checks; ``size_of(shape, para, run)`` returns a run's font size in points."""
    findings = []
    findings += _check_capacity_and_samples(prs, manifest, Finding)
    findings += _check_repetition(prs, manifest, Finding, slide_w, slide_h)
    findings += _check_cover(prs, Finding, size_of, slide_w, slide_h)
    return findings
