#!/usr/bin/env python
"""Build a deck from a built-in template: clone prototype pages, fill named slots.

Use from a build script in the code interpreter::

    import sys
    sys.path.insert(0, "skills/bisheng-pptx/scripts")
    from deck_template import Deck

    deck = Deck.open("gov-red", footer="某某集团 · 2026 年度工作汇报")
    s = deck.add("cover")
    deck.fill(s, "title", "2026 年度工作汇报")
    deck.drop(s, "kicker")
    s = deck.add("chart-insight")
    deck.fill(s, "title", "营收连续四年增长，2026 年预计突破 130 亿元")
    deck.chart(s, "chart", "column", ["2022", "2023", "2024", "2025"], [("营收（亿元）", [82, 95, 108, 121])])
    deck.fill(s, "insight.body", ["要点一", "要点二"])
    deck.save("output/年度工作汇报.pptx")

``python skills/bisheng-pptx/scripts/list_templates.py <id>`` prints every layout
and slot. ``save`` drops the prototype pages and writes
``scratch/<name>.manifest.json``, which ``inspect_deck.py`` reads for the capacity
and leftover-sample checks.
"""

from __future__ import annotations

import copy
import json
import os
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt

_HERE = os.path.dirname(os.path.abspath(__file__))
if _HERE not in sys.path:
    sys.path.insert(0, _HERE)

from deck_render import add_chart, build_table  # noqa: E402

__all__ = ["TEMPLATES_DIR", "Deck", "text_units"]

TEMPLATES_DIR = os.path.normpath(os.path.join(_HERE, "..", "templates"))
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_SKIP_RELS = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.CHART}
_FOOT_LIMIT_PT = 492.0  # tables must end above the footer line
_TEXT_ROLES = {"title", "subtitle", "body", "caption", "label", "number", "footer"}


def text_units(text: str) -> float:
    """字数: CJK / full-width counts 1, anything else 0.5 — the unit of every capacity."""
    return sum(1.0 if ord(ch) > 0x2E7F else 0.5 for ch in text if ch not in "\r\n")


def _load_layouts(template_id: str) -> dict:
    path = os.path.join(TEMPLATES_DIR, template_id, "layouts.json")
    if not os.path.exists(path):
        known = sorted(d for d in os.listdir(TEMPLATES_DIR) if os.path.isdir(os.path.join(TEMPLATES_DIR, d)))
        raise ValueError(f"未知模板 {template_id!r}；可用模板：{', '.join(known)}")
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _has_chart(element) -> bool:
    return element.find(f".//{{{_C_NS}}}chart") is not None


class Deck:
    """A deck under construction on top of one template."""

    def __init__(self, prs, spec: dict, footer: str | None):
        self.prs = prs
        self.spec = spec
        self.footer = footer
        self.layouts = {item["layout_id"]: item for item in spec["layouts"]}
        self._prototype_count = len(prs.slides)
        self._pages: list[dict] = []  # parallel to slides appended after the prototypes

    # ---- construction -------------------------------------------------------
    @classmethod
    def open(cls, template_id: str, footer: str | None = None) -> Deck:
        """Open a built-in template. ``footer`` fills every page footer; ``None`` removes it."""
        spec = _load_layouts(template_id)
        prs = Presentation(os.path.join(TEMPLATES_DIR, template_id, "template.pptx"))
        return cls(prs, spec, footer)

    def add(self, layout_id: str):
        """Append a copy of prototype ``layout_id`` and return the new slide."""
        if layout_id not in self.layouts:
            raise ValueError(f"模板 {self.spec['id']} 没有版式 {layout_id!r}；可用：{', '.join(self.layouts)}")
        layout = self.layouts[layout_id]
        src = self.prs.slides[layout["slide_index"]]
        new = self.prs.slides.add_slide(src.slide_layout)
        for shape in list(new.shapes):
            shape._element.getparent().remove(shape._element)

        rid_map = {}
        for rid, rel in src.part.rels.items():
            if rel.reltype in _SKIP_RELS:
                continue
            if rel.is_external:
                rid_map[rid] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            else:
                rid_map[rid] = new.part.rels.get_or_add(rel.reltype, rel.target_part)

        src_bg = src._element.cSld.find(qn("p:bg"))
        if src_bg is not None:
            new._element.cSld.insert(0, copy.deepcopy(src_bg))
        for shape in src.shapes:
            if (
                shape.shape_type is not None
                and shape._element.tag == qn("p:graphicFrame")
                and _has_chart(shape._element)
            ):
                continue  # charts are re-created by Deck.chart (a cloned rel would share one chart part)
            el = copy.deepcopy(shape._element)
            for node in el.iter():
                for attr, value in list(node.attrib.items()):
                    if attr.startswith(f"{{{_R_NS}}}") and value in rid_map:
                        node.set(attr, rid_map[value])
            new.shapes._spTree.append(el)

        page = {"layout_id": layout_id, "filled": {}, "over_capacity": [], "dropped": [], "objects": {}}
        self._pages.append(page)
        if self._slot(layout, "footer", required=False):
            if self.footer:
                self.fill(new, "footer", self.footer)
            else:
                self.drop(new, "footer")
        return new

    # ---- slot helpers ---------------------------------------------------------
    def _page(self, slide) -> tuple[dict, dict]:
        index = self._slide_position(slide) - self._prototype_count
        if index < 0:
            raise ValueError("只能修改通过 deck.add() 新增的页面，不能改模板原型页")
        page = self._pages[index]
        return page, self.layouts[page["layout_id"]]

    def _slide_position(self, slide) -> int:
        for i, s in enumerate(self.prs.slides):
            if s.slide_id == slide.slide_id:
                return i
        raise ValueError("这个 slide 不属于当前 deck")

    @staticmethod
    def _slot(layout: dict, name: str, required: bool = True):
        for slot in layout["slots"]:
            if slot["name"] == name:
                return slot
        if required:
            names = [s["name"] for s in layout["slots"]]
            raise KeyError(f"版式 {layout['layout_id']} 没有插槽 {name!r}；可用插槽：{', '.join(names)}")
        return None

    @staticmethod
    def _shape(slide, name: str):
        for shape in slide.shapes:
            if shape.name == name:
                return shape
        return None

    # ---- text -----------------------------------------------------------------
    def fill(self, slide, slot: str, text) -> None:
        """Write text into a slot, keeping the template's font, size, colour and bullets.

        ``text`` is a string (``\\n`` starts a new paragraph) or a list of paragraphs.
        Exceeding the slot capacity prints a warning and is recorded for inspect_deck.
        """
        page, layout = self._page(slide)
        spec = self._slot(layout, slot)
        if spec["role"] not in _TEXT_ROLES:
            raise ValueError(f"插槽 {slot!r} 是 {spec['role']}，请用 deck.chart / deck.table / deck.image")
        shape = self._shape(slide, slot)
        if shape is None:
            raise ValueError(f"插槽 {slot!r} 已被 drop，不能再填写")
        lines = text.split("\n") if isinstance(text, str) else [str(x) for x in text]
        lines = lines or [""]
        _replace_paragraphs(shape, lines)
        used = text_units("".join(lines))
        page["filled"][slot] = used
        if used > spec["capacity"]:
            entry = {"slot": slot, "used": used, "capacity": spec["capacity"]}
            page["over_capacity"].append(entry)
            print(
                f"[超容量] 第{self._pages.index(page) + 1}页 {slot}: {used:g} 字 > 容量 {spec['capacity']} 字 → 精简文案或拆页"
            )

    def drop(self, slide, name: str) -> None:
        """Remove a slot or a whole group (``card3`` removes ``card3.*``) you don't need."""
        page, layout = self._page(slide)
        removed = 0
        for shape in list(slide.shapes):
            if shape.name == name or shape.name.startswith(name + "."):
                shape._element.getparent().remove(shape._element)
                removed += 1
        if not removed and not any(s["name"] == name for s in layout["slots"]) and name not in layout["groups"]:
            raise KeyError(f"版式 {layout['layout_id']} 没有插槽或分组 {name!r}")
        page["dropped"].append(name)

    def notes(self, slide, text: str) -> None:
        """Speaker notes for a page."""
        slide.notes_slide.notes_text_frame.text = text

    # ---- native objects -----------------------------------------------------------
    def chart(self, slide, slot: str, kind: str, categories, series, number_format: str = "0") -> None:
        """Native chart in a ``chart_area`` slot.

        kind: column | bar | line | stacked | pie | doughnut | waterfall.
        series: list of (label, values). waterfall: one series, first/last are totals.
        """
        page, layout = self._page(slide)
        spec = self._slot(layout, slot)
        if spec["role"] != "chart_area":
            raise ValueError(f"插槽 {slot!r} 不是图表区")
        x, y, w, h = spec["bounds"]
        frame = add_chart(slide.shapes, x, y, w, h, kind, categories, series, spec["chart_style"], number_format)
        frame.name = slot
        page["objects"][slot] = kind

    def table(self, slide, slot: str, header, rows, col_widths=None, highlight_rows=()) -> None:
        """Native table in a ``table`` slot; ``highlight_rows`` are 0-based data-row indexes to flag."""
        page, layout = self._page(slide)
        spec = self._slot(layout, slot)
        if spec["role"] != "table":
            raise ValueError(f"插槽 {slot!r} 不是表格")
        old = self._shape(slide, slot)
        if old is not None:
            old._element.getparent().remove(old._element)
        x, y, w, _ = spec["bounds"]
        row_h = float(spec.get("row_h", 34.0))
        n = len(rows) + 1
        room = _FOOT_LIMIT_PT - y
        if n * row_h > room:
            row_h = max(26.0, room / n)
        if n * row_h > room:
            page["over_capacity"].append({"slot": slot, "used": n, "capacity": int(room // 26)})
            print(
                f"[超容量] 第{self._pages.index(page) + 1}页 {slot}: {n} 行放不下（最多约 {int(room // 26)} 行）→ 拆成两页"
            )
        frame = build_table(
            slide.shapes,
            x,
            y,
            w,
            header,
            rows,
            spec["style"],
            col_widths=col_widths or spec.get("col_widths"),
            row_h=row_h,
            highlight_rows=highlight_rows,
        )
        frame.name = slot
        note = self._shape(slide, "note")
        if note is not None:
            note.top = Pt(y + n * row_h + 16)
        page["objects"][slot] = "table"

    def image(self, slide, slot: str, path: str, fit: str = "cover") -> None:
        """Place a picture in an ``image_area`` slot. fit=cover crops to fill, contain letterboxes."""
        from PIL import Image

        page, layout = self._page(slide)
        spec = self._slot(layout, slot)
        if spec["role"] != "image_area":
            raise ValueError(f"插槽 {slot!r} 不是图片区")
        x, y, w, h = spec["bounds"]
        with Image.open(path) as im:
            iw, ih = im.size
        box_ratio, img_ratio = w / h, iw / ih
        if fit == "contain":
            if img_ratio > box_ratio:
                pw, ph = w, w / img_ratio
            else:
                pw, ph = h * img_ratio, h
            pic = slide.shapes.add_picture(path, Pt(x + (w - pw) / 2), Pt(y + (h - ph) / 2), Pt(pw), Pt(ph))
        else:
            pic = slide.shapes.add_picture(path, Pt(x), Pt(y), Pt(w), Pt(h))
            if img_ratio > box_ratio:
                excess = 1 - box_ratio / img_ratio
                pic.crop_left = pic.crop_right = excess / 2
            else:
                excess = 1 - img_ratio / box_ratio
                pic.crop_top = pic.crop_bottom = excess / 2
        for shape in list(slide.shapes):
            if shape.name.startswith(slot + "."):
                shape._element.getparent().remove(shape._element)
        pic.name = slot
        page["objects"][slot] = "image"

    # ---- output -------------------------------------------------------------------
    def save(self, path: str, manifest_dir: str = "scratch") -> str:
        """Drop the prototype pages, save the deck and write ``<manifest_dir>/<name>.manifest.json``."""
        if not self._pages:
            raise ValueError("还没有通过 deck.add() 添加任何页面")
        slide_ids = self.prs.slides._sldIdLst
        for entry in list(slide_ids)[: self._prototype_count]:
            self.prs.part.drop_rel(entry.get(qn("r:id")))
            slide_ids.remove(entry)
        self._prototype_count = 0
        folder = os.path.dirname(path)
        if folder:
            os.makedirs(folder, exist_ok=True)
        self.prs.save(path)

        pages = []
        for i, page in enumerate(self._pages, start=1):
            layout = self.layouts[page["layout_id"]]
            unfilled = []
            for spec in layout["slots"]:
                name = spec["name"]
                gone = any(name == d or name.startswith(d + ".") for d in page["dropped"])
                if gone or name in page["filled"] or name in page["objects"]:
                    continue
                unfilled.append(name)
            pages.append(
                {
                    "index": i,
                    "layout_id": page["layout_id"],
                    "filled": page["filled"],
                    "over_capacity": page["over_capacity"],
                    "unfilled": unfilled,
                    "objects": page["objects"],
                }
            )
        manifest = {"template": self.spec["id"], "deck": os.path.basename(path), "pages": pages}
        os.makedirs(manifest_dir, exist_ok=True)
        stem = os.path.splitext(os.path.basename(path))[0]
        manifest_path = os.path.join(manifest_dir, f"{stem}.manifest.json")
        with open(manifest_path, "w", encoding="utf-8") as fh:
            json.dump(manifest, fh, ensure_ascii=False, indent=1)
        print(f"[saved] {path}（{len(pages)} 页，模板 {self.spec['id']}）；清单 {manifest_path}")
        return manifest_path


def _replace_paragraphs(shape, lines: list[str]) -> None:
    """Rewrite a text frame: line 1 keeps paragraph 1's format, later lines paragraph 2's (or 1's)."""
    body = shape.text_frame._txBody
    paras = body.findall(qn("a:p"))
    first = paras[0]
    rest_proto = paras[1] if len(paras) > 1 else paras[0]
    proto_first = copy.deepcopy(first)
    proto_rest = copy.deepcopy(rest_proto)
    for p in paras:
        body.remove(p)
    for i, line in enumerate(lines):
        p = copy.deepcopy(proto_first if i == 0 else proto_rest)
        runs = p.findall(qn("a:r"))
        for extra in runs[1:]:
            p.remove(extra)
        for fld in p.findall(qn("a:fld")):
            p.remove(fld)
        if runs:
            run = runs[0]
        else:
            run = p.makeelement(qn("a:r"), {})
            end = p.find(qn("a:endParaRPr"))
            if end is not None:
                end.addprevious(run)
            else:
                p.append(run)
        t = run.find(qn("a:t"))
        if t is None:
            t = run.makeelement(qn("a:t"), {})
            run.append(t)
        t.text = line
        if line != line.strip():
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        body.append(p)
