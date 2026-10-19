"""解析云学堂在线课堂导出表.

第三方课程导入和课程目录导入都读这一份表. 表头与导出文件一致, 只取课程和目录要用的列.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO

import openpyxl

from bisheng.common.errcode.portal_course import PortalCourseImportError
from bisheng.shougang_portal_course.domain.schemas.portal_course_schema import (
    CourseImportIssue,
    CourseTag,
    validate_media_url,
)
from bisheng.shougang_portal_course.domain.services.catalog_service import MAX_CATALOG_DEPTH

TEMPLATE_HEADERS = (
    "课堂名称",
    "课件名称",
    "全路径",
    "分类名称",
    "上传者",
    "贡献者",
    "账号",
    "作者",
    "账号",
    "主讲老师",
    "账号",
    "辅讲老师",
    "账号",
    "课件类型",
    "课件状态",
    "标准时长",
    "学分",
    "平台课程编号",
    "分享链接",
    "课程简介",
    "推荐系数",
    "价格",
    "编码",
    "版本",
    "标签",
    "是否上架",
    "是否隐藏",
    "是否锁定",
    "是否原创",
    "有效性",
    "过期时间",
    "是否支持App",
    "上传日期",
    "更新日期",
    "版本号",
    "浏览权限设置",
    "浏览权限",
    "下载设置",
    "下载权限",
    "app缓存权限设置",
    "app缓存权限权限",
    "来源",
    "环境",
    "浏览人数",
    "学习人数",
    "点赞人数",
    "评论人数",
    "下载次数",
    "综合评分",
    "是否被课程引用",
)
REQUIRED_HEADERS = (
    "课件名称",
    "全路径",
    "分类名称",
    "主讲老师",
    "标准时长",
    "平台课程编号",
    "分享链接",
    "课程简介",
    "更新日期",
    "标签",
)
SHEET_NAME = "Sheet0"
MAX_IMPORT_ROWS = 2000
MAX_NAME_LENGTH = 200
MAX_INSTRUCTOR_LENGTH = 100
MAX_EXTERNAL_ID_LENGTH = 128
MAX_CATALOG_NAME_LENGTH = 200
MAX_TAG_LENGTH = 50
_EMPTY_TOKENS = {"--", "-"}
_TAG_SPLIT = re.compile("[,\\uFF0C;\\uFF1B\\u3001|]+")
_DURATION = re.compile(r"^(?:(\d+)时)?(?:(\d+)分)?(?:(\d+)秒)?$")
_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y/%m/%d %H:%M",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%Y.%m.%d",
)


@dataclass
class ParsedCourseRow:
    """一行已经通过校验的在线课堂课程."""

    row_idx: int
    external_id: str
    name: str
    instructor: str
    catalog_names: list[str]
    source_updated_at: datetime | None
    description: str
    tags: list[CourseTag]
    external_url: str
    duration_seconds: int


def build_template() -> bytes:
    """生成只含在线课堂表头的空 xlsx.

    Returns:
        工作表 Sheet0 的表头与云学堂导出一致, 没有数据行.
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    for col, header in enumerate(TEMPLATE_HEADERS, start=1):
        sheet.cell(row=1, column=col, value=header)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def parse_workbook(content: bytes) -> tuple[list[ParsedCourseRow], list[CourseImportIssue], int]:
    """解析在线课堂导出表中的课程行.

    Args:
        content: xlsx 文件字节.

    Returns:
        通过校验的行, 行级问题, 以及非空数据行数.

    Raises:
        PortalCourseImportError: 文件无法打开, 缺少必需表头, 或数据行超过 2000.
    """
    headers, data_rows = _load_sheet(content)
    header_index = _header_index(headers)
    parsed: list[ParsedCourseRow] = []
    issues: list[CourseImportIssue] = []
    total = 0
    for offset, raw in enumerate(data_rows, start=2):
        if _is_empty_row(raw):
            continue
        total += 1
        if total > MAX_IMPORT_ROWS:
            raise PortalCourseImportError(msg=f"导入行数不能超过 {MAX_IMPORT_ROWS}")
        try:
            parsed.append(_parse_row(offset, raw, header_index))
        except ValueError as exc:
            issues.append(
                CourseImportIssue(
                    row=offset,
                    code="invalid_row",
                    message=str(exc),
                    recoverable=False,
                )
            )
    return parsed, issues, total


def catalog_node_count(rows: list[ParsedCourseRow]) -> int:
    """统计这些课程行会确保的目录节点数.

    Args:
        rows: 已通过校验的课程行.

    Returns:
        从根到叶每一级去重后的目录个数.
    """
    nodes: set[tuple[str, ...]] = set()
    for row in rows:
        acc: list[str] = []
        for name in row.catalog_names:
            acc.append(name)
            nodes.add(tuple(acc))
    return len(nodes)


def _load_sheet(content: bytes) -> tuple[list[str], list[tuple]]:
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=True)
    except Exception as exc:
        raise PortalCourseImportError(msg="无法解析 Excel 文件") from exc
    try:
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [_cell_text(cell) for cell in rows[0]]
            if all(name in headers for name in REQUIRED_HEADERS):
                return headers, [tuple(row or ()) for row in rows[1:]]
    finally:
        workbook.close()
    required = "、".join(REQUIRED_HEADERS)
    raise PortalCourseImportError(msg=f"表头必须包含 {required}")


def _header_index(headers: list[str]) -> dict[str, int]:
    index: dict[str, int] = {}
    for position, name in enumerate(headers):
        index.setdefault(name, position)
    return index


def _parse_row(row_idx: int, raw: tuple, header_index: dict[str, int]) -> ParsedCourseRow:
    def cell(name: str) -> object:
        position = header_index[name]
        if position >= len(raw):
            return None
        return raw[position]

    name = _required_text(cell("课件名称"), field="课件名称", row_idx=row_idx, max_length=MAX_NAME_LENGTH)
    external_id = _required_text(
        cell("平台课程编号"),
        field="平台课程编号",
        row_idx=row_idx,
        max_length=MAX_EXTERNAL_ID_LENGTH,
    )
    instructor = _limited_text(
        _blank_token(_cell_text(cell("主讲老师"))),
        field="主讲老师",
        row_idx=row_idx,
        max_length=MAX_INSTRUCTOR_LENGTH,
    )
    catalog_names = _catalog_names(row_idx, cell("全路径"), cell("分类名称"))
    description = _cell_text(cell("课程简介"))
    source_updated_at = _parse_date(row_idx, cell("更新日期"))
    tags = _parse_tags(row_idx, cell("标签"))
    external_url = _parse_url(row_idx, cell("分享链接"))
    duration_seconds = _parse_duration(row_idx, cell("标准时长"))
    return ParsedCourseRow(
        row_idx=row_idx,
        external_id=external_id,
        name=name,
        instructor=instructor,
        catalog_names=catalog_names,
        source_updated_at=source_updated_at,
        description=description,
        tags=tags,
        external_url=external_url,
        duration_seconds=duration_seconds,
    )


def _catalog_names(row_idx: int, path_value: object, category_value: object) -> list[str]:
    path_text = _cell_text(path_value)
    category = _cell_text(category_value)
    if not path_text:
        raise ValueError(f"第 {row_idx} 行: 全路径不能为空")
    if not category:
        raise ValueError(f"第 {row_idx} 行: 分类名称不能为空")
    parts = [part.strip() for part in path_text.split("->")]
    if any(not part for part in parts):
        raise ValueError(f"第 {row_idx} 行: 全路径格式无效")
    if parts[-1] != category:
        raise ValueError(f"第 {row_idx} 行: 分类名称必须与全路径最后一级一致")
    if len(parts) > MAX_CATALOG_DEPTH:
        raise ValueError(f"第 {row_idx} 行: 目录层级不能超过 {MAX_CATALOG_DEPTH} 级")
    for part in parts:
        if len(part) > MAX_CATALOG_NAME_LENGTH:
            raise ValueError(f"第 {row_idx} 行: 目录名称不能超过 {MAX_CATALOG_NAME_LENGTH} 字")
        if "/" in part:
            raise ValueError(f"第 {row_idx} 行: 目录名称不能包含 /")
    return parts


def _parse_duration(row_idx: int, value: object) -> int:
    text = _cell_text(value)
    matched = _DURATION.fullmatch(text)
    if matched is None or not any(matched.groups()):
        raise ValueError(f"第 {row_idx} 行: 标准时长无法解析")
    hours = int(matched.group(1) or 0)
    minutes = int(matched.group(2) or 0)
    seconds = int(matched.group(3) or 0)
    return hours * 3600 + minutes * 60 + seconds


def _parse_url(row_idx: int, value: object) -> str:
    text = _blank_token(_cell_text(value))
    if not text:
        return ""
    try:
        return validate_media_url(text)
    except ValueError as exc:
        raise ValueError(f"第 {row_idx} 行: 分享链接必须是有效的 http(s) 链接") from exc


def _parse_date(row_idx: int, value: object) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = _cell_text(value)
    if not text or text in _EMPTY_TOKENS:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    raise ValueError(f"第 {row_idx} 行: 更新日期格式无效")


def _parse_tags(row_idx: int, value: object) -> list[CourseTag]:
    text = _cell_text(value)
    if not text or text in _EMPTY_TOKENS:
        return []
    tags: list[CourseTag] = []
    seen: set[str] = set()
    for chunk in _TAG_SPLIT.split(text):
        label = _blank_token(chunk.strip())
        if not label or label in seen:
            continue
        if len(label) > MAX_TAG_LENGTH:
            raise ValueError(f"第 {row_idx} 行: 标签不能超过 {MAX_TAG_LENGTH} 字")
        seen.add(label)
        tags.append(CourseTag(label=label, display_type="gray"))
    return tags


def _required_text(value: object, *, field: str, row_idx: int, max_length: int) -> str:
    text = _cell_text(value)
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    if not text:
        raise ValueError(f"第 {row_idx} 行: {field}不能为空")
    return _limited_text(text, field=field, row_idx=row_idx, max_length=max_length)


def _limited_text(text: str, *, field: str, row_idx: int, max_length: int) -> str:
    if len(text) > max_length:
        raise ValueError(f"第 {row_idx} 行: {field}不能超过 {max_length} 字")
    return text


def _blank_token(text: str) -> str:
    # 云学堂用 -- 和 - 表示这一格没有值.
    return "" if text in _EMPTY_TOKENS else text


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def _is_empty_row(row: tuple | None) -> bool:
    if not row:
        return True
    return all(cell is None or str(cell).strip() == "" for cell in row)
