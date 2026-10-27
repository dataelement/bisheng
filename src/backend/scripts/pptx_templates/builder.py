"""Low-level drawing for template generation: shapes, slots, theme XML.

Every text box that the model is expected to rewrite is a *slot*: its shape name
is recorded together with a role, a capacity (in 字数, see ``text_units``) and the
sample text, so the shipped ``layouts.json`` can drive ``Deck.fill`` and the
capacity / leftover-sample checks. Names of the form ``group.part`` belong to a
droppable group (``card3.title`` → group ``card3``).
"""

from __future__ import annotations

import math
import os
import sys
import uuid

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Pt
from themes import Theme

_SKILL_SCRIPTS = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "..", "..", "bisheng", "linsight", "builtin_skills", "bisheng-pptx",
                 "scripts")
)
if _SKILL_SCRIPTS not in sys.path:
    sys.path.insert(0, _SKILL_SCRIPTS)

from deck_render import add_chart, build_table, set_run_font  # noqa: E402

SLIDE_W = 960.0
SLIDE_H = 540.0
CAPACITY_SLACK = 0.9  # leave ~10% for font-metric differences on the opening machine
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def text_units(text: str) -> float:
    """字数 used for every capacity: CJK / full-width counts 1, anything else 0.5."""
    units = 0.0
    for ch in text:
        if ch in "\n\r":
            continue
        units += 1.0 if ord(ch) > 0x2E7F else 0.5
    return units


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_alpha(shape, alpha: float) -> None:
    """Fill transparency: python-pptx has no API, write <a:alpha> into the solid fill."""
    srgb = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    srgb.append(srgb.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))}))


class SlideBuilder:
    """Draws one prototype slide and records its slots."""

    def __init__(self, prs, theme: Theme, layout_id: str, family: str, name: str, use_for: str):
        self.theme = theme
        self.layout_id = layout_id
        self.family = family
        self.name = name
        self.use_for = use_for
        blank = prs.slide_layouts[0]
        self.slide = prs.slides.add_slide(blank)
        self.slots: list[dict] = []
        self.samples: list[str] = []
        self.groups: dict[str, list[str]] = {}
        self.extra: dict = {}
        self._names: set[str] = set()

    # ---- bookkeeping -------------------------------------------------------
    def _register(self, shape, name: str) -> None:
        if name in self._names:
            raise ValueError(f"{self.theme.id}/{self.layout_id}: duplicate shape name {name}")
        self._names.add(name)
        shape.name = name
        if "." in name:
            group = name.split(".", 1)[0]
            self.groups.setdefault(group, []).append(name)

    # ---- canvas ------------------------------------------------------------
    def background(self, color: str) -> None:
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    # ---- shapes ------------------------------------------------------------
    def rect(self, name, x, y, w, h, fill=None, line=None, line_w=0.75, radius=0.0, alpha=None, shape=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        shp = self.slide.shapes.add_shape(kind, Pt(x), Pt(y), Pt(w), Pt(h))
        self._register(shp, name)
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            shp.adjustments[0] = radius or 0.08
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
            if alpha is not None:
                set_alpha(shp, alpha)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(line_w)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        if shp.has_text_frame:
            shp.text_frame.text = ""
        return shp

    def oval(self, name, x, y, w, h, fill=None, line=None, line_w=0.75, alpha=None):
        return self.rect(name, x, y, w, h, fill=fill, line=line, line_w=line_w, alpha=alpha, shape=MSO_SHAPE.OVAL)

    def hline(self, name, x, y, w, color, weight=0.75):
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x + w), Pt(y))
        self._register(conn, name)
        conn.line.color.rgb = rgb(color)
        conn.line.width = Pt(weight)
        return conn

    def vline(self, name, x, y, h, color, weight=0.75):
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(x), Pt(y), Pt(x), Pt(y + h))
        self._register(conn, name)
        conn.line.color.rgb = rgb(color)
        conn.line.width = Pt(weight)
        return conn

    # ---- text --------------------------------------------------------------
    def text(
        self,
        name,
        x,
        y,
        w,
        h,
        sample,
        *,
        size,
        color,
        bold=False,
        font="body",
        align="left",
        anchor="top",
        role="body",
        line_spacing=1.2,
        slot=True,
        bullets=False,
        residual=None,
    ):
        """Add a text box. ``sample`` may be a list (one paragraph per item)."""
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        self._register(box, name)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = _ANCHOR[anchor]
        lines = sample if isinstance(sample, list) else [sample]
        ea = self.theme.title_font if font == "title" else self.theme.body_font
        latin = self.theme.latin_font
        for i, line in enumerate(lines):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = _ALIGN[align]
            para.line_spacing = line_spacing
            if bullets:
                _add_bullet(para, color=self.theme.accent if not self.theme.all_dark else self.theme.primary)
            if i > 0 and not bullets:
                para.space_before = Pt(size * 0.3)
            elif i > 0:
                para.space_before = Pt(size * 0.35)
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = rgb(color)
            set_run_font(run, latin, ea)
        text = "\n".join(lines)
        if slot if residual is None else residual:
            self.samples.append(text)
        if slot:
            indent = size * 1.1 if bullets else 0.0
            per_line = max(1, math.floor((w - indent) / size))
            max_lines = max(1, math.floor(h / (size * line_spacing * 1.08)))
            capacity = max(1, math.floor(per_line * max_lines * CAPACITY_SLACK))
            used = text_units(text)
            if used > capacity:
                raise ValueError(f"{self.theme.id}/{self.layout_id}/{name}: sample uses {used} > capacity {capacity}")
            self.slots.append(
                {
                    "name": name,
                    "role": role,
                    "capacity": capacity,
                    "max_lines": max_lines,
                    "font_size": size,
                    "sample": text,
                    "bullets": bullets,
                }
            )
        return box

    def page_number(self, name, x, y, w, h, size, color, align="right"):
        """Slide-number field; not a slot."""
        box = self.slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        self._register(box, name)
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        para = tf.paragraphs[0]
        para.alignment = _ALIGN[align]
        fld = etree.SubElement(para._p, qn("a:fld"), {"id": "{" + str(uuid.uuid4()).upper() + "}", "type": "slidenum"})
        rpr = etree.SubElement(fld, qn("a:rPr"), {"lang": "zh-CN", "sz": str(int(size * 100))})
        fill = etree.SubElement(rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), {"val": color})
        etree.SubElement(rpr, qn("a:latin"), {"typeface": self.theme.latin_font})
        t = etree.SubElement(fld, qn("a:t"))
        t.text = "‹#›"
        return box

    # ---- native objects ----------------------------------------------------
    def chart(self, name, x, y, w, h, categories, series, kind="column", number_format="0"):
        """Sample native chart; ``Deck.chart`` re-creates it in the slot with real data."""
        style = chart_style(self.theme)
        frame = add_chart(self.slide.shapes, x, y, w, h, kind, categories, series, style, number_format)
        self._register(frame, name)
        self.slots.append({"name": name, "role": "chart_area", "bounds": [x, y, w, h], "chart_style": style})
        return frame

    def table(self, name, x, y, w, header, rows, col_widths=None, row_h=34.0, highlight_rows=()):
        """Sample native table; ``Deck.table`` rebuilds it from the stored style."""
        style = table_style(self.theme)
        frame = build_table(self.slide.shapes, x, y, w, header, rows, style, col_widths=col_widths, row_h=row_h,
                            highlight_rows=highlight_rows)
        self._register(frame, name)
        self.slots.append({
            "name": name, "role": "table", "bounds": [x, y, w, row_h * (len(rows) + 1)], "row_h": row_h,
            "col_widths": col_widths, "style": style,
        })
        for cell_text in header + [c for r in rows for c in r]:
            self.samples.append(cell_text)
        return frame

    def image_area(self, name, x, y, w, h, label="示例：图片区，替换为现场照片或产品截图"):
        """Placeholder for a picture: a tinted box plus a label, both in group ``name``."""
        t = self.theme
        self.rect(f"{name}.box", x, y, w, h, fill=t.surface, radius=t.card_radius)
        self.text(
            f"{name}.label",
            x + 24,
            y + h / 2 - 12,
            w - 48,
            24,
            label,
            size=12,
            color=t.muted,
            align="center",
            role="caption",
            slot=False,
            residual=True,
        )
        self.slots.append({"name": name, "role": "image_area", "bounds": [x, y, w, h]})

    # ---- export ------------------------------------------------------------
    def notes(self, text: str) -> None:
        self.slide.notes_slide.notes_text_frame.text = text

    def describe(self, index: int) -> dict:
        return {
            "layout_id": self.layout_id,
            "slide_index": index,
            "family": self.family,
            "name": self.name,
            "use_for": self.use_for,
            "slots": self.slots,
            "groups": {g: names for g, names in self.groups.items() if not g.startswith("deco")},
            "samples": [s for s in self.samples if s.strip()],
        }


# ---------------------------------------------------------------------------
# Shared styling — the same rules Deck.chart / Deck.table apply at build time.
# ---------------------------------------------------------------------------


def table_style(theme: Theme) -> dict:
    dark = theme.all_dark
    return {
        "font": theme.body_font,
        "latin": theme.latin_font,
        "header_fill": theme.primary if not dark else theme.secondary,
        "header_color": "FFFFFF" if not dark else theme.text,
        "header_size": 13,
        "body_size": 12,
        "body_color": theme.text,
        "stripe_fill": theme.surface,
        "body_fill": theme.bg,
        "line_color": theme.line,
        "highlight_fill": "FBE9E7" if not dark else "3A2A33",
        "highlight_color": "B3261E" if not dark else "FF8A80",
    }


def chart_style(theme: Theme) -> dict:
    if theme.all_dark:
        palette = [theme.primary, theme.secondary, "F2B84B", theme.muted, "8C6CE6"]
    else:
        palette = [theme.primary, theme.accent, theme.secondary, theme.muted, theme.line]
    return {
        "palette": palette,
        "text": theme.text,
        "muted": theme.muted,
        "line": theme.line,
        "font": theme.latin_font,
        "negative": "C0504D" if not theme.all_dark else "FF8A80",
    }


def _add_bullet(para, color: str) -> None:
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(14))))
    pPr.set("indent", str(-int(Pt(14))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr", "a:buFont"):
        for old in pPr.findall(qn(tag)):
            pPr.remove(old)
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr"), {"val": color})
    etree.SubElement(pPr, qn("a:buFont"), {"typeface": "Arial"})
    etree.SubElement(pPr, qn("a:buChar"), {"char": "•"})


# ---------------------------------------------------------------------------
# Presentation-level setup
# ---------------------------------------------------------------------------


def prepare_presentation(prs, theme: Theme) -> None:
    """16:9 canvas, a single blank layout, and theme colors/fonts from the tokens."""
    prs.slide_width = Pt(SLIDE_W)
    prs.slide_height = Pt(SLIDE_H)
    master = prs.slide_masters[0]
    layouts = master.slide_layouts
    blank = next(l for l in layouts if l.name == "Blank")
    for layout in list(layouts):
        if layout is not blank:
            layouts.remove(layout)
    for shape in list(blank.placeholders):
        shape._element.getparent().remove(shape._element)
    for shape in list(master.placeholders):
        shape._element.getparent().remove(shape._element)
    _patch_theme(master.part.part_related_by(RT.THEME), theme)


def _patch_theme(theme_part, theme: Theme) -> None:
    root = etree.fromstring(theme_part.blob)
    ns = {"a": A_NS}
    root.set("name", f"{theme.id}")
    scheme = root.find(".//a:clrScheme", ns)
    scheme.set("name", theme.id)
    dark_text = theme.text if not theme.all_dark else theme.dark_bg
    light_bg = theme.bg if not theme.all_dark else theme.text
    mapping = {
        "dk1": dark_text,
        "lt1": light_bg if not theme.all_dark else "FFFFFF",
        "dk2": theme.dark_bg,
        "lt2": theme.surface if not theme.all_dark else "F2F5FA",
        "accent1": theme.primary,
        "accent2": theme.accent,
        "accent3": theme.secondary,
        "accent4": theme.muted,
        "accent5": theme.on_dark_muted,
        "accent6": theme.line,
        "hlink": theme.secondary,
        "folHlink": theme.muted,
    }
    for tag, color in mapping.items():
        el = scheme.find(f"a:{tag}", ns)
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, qn("a:srgbClr"), {"val": color})
    font_scheme = root.find(".//a:fontScheme", ns)
    font_scheme.set("name", theme.id)
    for tag, ea in (("majorFont", theme.title_font), ("minorFont", theme.body_font)):
        node = font_scheme.find(f"a:{tag}", ns)
        node.find("a:latin", ns).set("typeface", theme.latin_font)
        node.find("a:ea", ns).set("typeface", ea)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
